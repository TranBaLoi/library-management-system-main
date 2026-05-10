# -*- coding: utf-8 -*-
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

src = r"C:\Users\Loi\Desktop\SQA\UnitTestReport_LibraryMS.xlsx"
out = r"C:\Users\Loi\Desktop\SQA\UnitTestReport_LibraryMS_Postman.xlsx"
wb = load_workbook(src)

thin_border = Border(left=Side(style="thin"), right=Side(style="thin"), top=Side(style="thin"), bottom=Side(style="thin"))
header_fill = PatternFill(start_color="D6E4F0", end_color="D6E4F0", fill_type="solid")
postman_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
section_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
title_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
cell_align = Alignment(vertical="top", wrap_text=True)
center_align = Alignment(horizontal="center", vertical="top", wrap_text=True)

base_url = "http://localhost:3000"

postman_auth = {
    "TC_AUTH_REG_01": "POST {{baseUrl}}/api/auth/register\nBody JSON: {fullName,email,password,confirmPassword,...}\n\nCheckDB SQL:\nSELECT id,email,fullName,role,status,isDeleted FROM User WHERE email='nguyenvana_postman@example.com';\nKỳ vọng: có 1 dòng user ACTIVE, role READER.",
    "TC_AUTH_REG_02": "POST /api/auth/register với email sai format\nKỳ vọng HTTP 400\n\nCheckDB SQL:\nSELECT COUNT(*) FROM User WHERE email='invalid-email';\nKỳ vọng: 0.",
    "TC_AUTH_REG_03": "POST /api/auth/register với password != confirm\nKỳ vọng HTTP 400\n\nCheckDB SQL: không có user mới cho email test.",
    "TC_AUTH_REG_04": "POST lại cùng email đã đăng ký\nKỳ vọng HTTP 409\n\nCheckDB SQL:\nSELECT COUNT(*) FROM User WHERE email='nguyenvana_postman@example.com';\nKỳ vọng: =1.",
    "TC_AUTH_REG_05": "POST register với password yếu\nKỳ vọng 400\nCheckDB: không tạo user.",
    "TC_AUTH_REG_06": "POST register fullName rỗng\nKỳ vọng 400\nCheckDB: không tạo user.",
    "TC_AUTH_REG_07": "POST register có phoneNumber hợp lệ\nCheckDB: phoneNumber lưu đúng.",
    "TC_AUTH_REG_08": "POST register phone không hợp lệ\nKỳ vọng 400\nCheckDB: không tạo user.",
    "TC_AUTH_LOGIN_01": "POST /api/auth/login\nBody email/password đúng\n\nCheckDB SQL:\nSELECT * FROM RefreshToken WHERE userId=(SELECT id FROM User WHERE email='nguyenvana_postman@example.com');\nKỳ vọng: có token mới.",
    "TC_AUTH_LOGIN_02": "POST login email không tồn tại\nKỳ vọng 401\nCheckDB: số token không tăng.",
    "TC_AUTH_LOGIN_03": "POST login password sai\nKỳ vọng 401\nCheckDB: token không tăng.",
    "TC_AUTH_LOGIN_04": "Set user INACTIVE rồi login\nKỳ vọng 401\nRollback: set ACTIVE lại.",
    "TC_AUTH_LOGIN_05": "Set user isDeleted=true rồi login\nKỳ vọng 401\nRollback: set isDeleted=false.",
    "TC_AUTH_LOGIN_06": "POST login email sai format\nKỳ vọng 400\nCheckDB: không token mới.",
    "TC_AUTH_LOGIN_07": "POST login password rỗng\nKỳ vọng 400\nCheckDB: không token mới.",
    "TC_AUTH_LOGIN_08": "Set firstLoginAt=NULL trước login\nSau login checkDB firstLoginAt != NULL.",
    "TC_AUTH_LOGIN_09": "User đã có firstLoginAt\nSau login checkDB firstLoginAt ổn định.",
    "TC_AUTH_LOGIN_10": "POST login rememberMe=true\nCheckDB expiresAt khoảng +30 ngày.",
    "TC_AUTH_REFRESH_01": "POST /api/auth/refresh với cookie refreshToken\nKỳ vọng 200, accessToken mới.",
    "TC_AUTH_REFRESH_02": "POST refresh với token giả\nKỳ vọng 401\nCheckDB: không token mới.",
    "TC_AUTH_REFRESH_03": "Set expiresAt của token < NOW() rồi refresh\nCheckDB token bị xóa.",
    "TC_AUTH_REFRESH_04": "User inactive rồi refresh\nKỳ vọng 401\nRollback status ACTIVE.",
    "TC_AUTH_CHANGEPWD_01": "POST /api/auth/change-password\nCheckDB: hash password đổi + RefreshToken count=0.",
    "TC_AUTH_CHANGEPWD_02": "POST change-password current sai\nKỳ vọng 400\nCheckDB: hash không đổi.",
    "TC_AUTH_CHANGEPWD_03": "POST change-password confirm mismatch\nKỳ vọng 400\nCheckDB: hash không đổi.",
    "TC_AUTH_CHANGEPWD_04": "POST change-password new==old\nKỳ vọng 400\nCheckDB: hash không đổi.",
    "TC_AUTH_CHANGEPWD_05": "Không có endpoint nhập userId tùy ý để test trực tiếp\nDùng integration với token user bị xóa.",
    "TC_AUTH_LOGOUT_01": "POST /api/auth/logout\nCheckDB token hiện tại bị xóa.",
    "TC_AUTH_LOGOUT_02": "Kiểm qua luồng đổi mật khẩu (logoutAll nội bộ)\nCheckDB RefreshToken của user = 0.",
    "TC_AUTH_CLEANUP_01": "Service cron nội bộ\nCheckDB trước/sau: SELECT COUNT(*) FROM RefreshToken WHERE expiresAt < NOW();",
    "TC_AUTH_CLEANUP_02": "Khi không có token hết hạn\nKỳ vọng count=0.",
}

postman_book = {
    "TC_BOOK_TRANSFORM_01": "GET {{baseUrl}}/api/books?page=1&limit=10\nKỳ vọng có categories, count, rating trong response.",
    "TC_BOOK_TRANSFORM_02": "Chọn book chưa có review -> GET /api/books\nCheckDB Review count=0 => averageRating=0.",
    "TC_BOOK_TRANSFORM_03": "Book thiếu category/edition/item -> GET /api/books\nKỳ vọng categories=[], count=0.",
    "TC_BOOK_FILTER_01": "GET /api/books\nCheckDB: response không gồm isDeleted=true.",
    "TC_BOOK_FILTER_02": "GET /api/books?authorIds=1&authorIds=2\nCheckDB authorId chỉ thuộc tập [1,2].",
    "TC_BOOK_FILTER_03": "GET /api/books?categoryIds=5&categoryIds=6\nCheckDB relation BookCategory.",
    "TC_BOOK_FILTER_04": "GET /api/books?languageCodes=vi&languageCodes=en\nCheckDB language in ('vi','en').",
    "TC_BOOK_FILTER_05": "GET /api/books?publishYearFrom=2000&publishYearTo=2020\nCheckDB publishYear trong khoảng.",
    "TC_BOOK_FILTER_06": "GET /api/books?availableAt=ebook\nCheckDB bookId nằm trong BookEdition format EBOOK.",
    "TC_BOOK_FILTER_07": "GET /api/books?availableAt=book-copy\nCheckDB bookId nằm trong BookItem status AVAILABLE.",
    "TC_BOOK_FILTER_08": "GET /api/books?availableAt=ebook&availableAt=book-copy\nCheckDB bookId thỏa cả hai điều kiện.",
    "TC_BOOK_SORT_01": "GET /api/books\nKỳ vọng createdAt desc.",
    "TC_BOOK_SORT_02": "GET /api/books?sortBy=title&sortOrder=asc\nCheckDB ORDER BY title ASC.",
    "TC_BOOK_SORT_03": "GET /api/books?sortBy=price&sortOrder=desc\nCheckDB ORDER BY price DESC.",
    "TC_BOOK_SORT_04": "GET /api/books?sortBy=unknownField&sortOrder=asc\nKỳ vọng fallback createdAt desc.",
    "TC_BOOK_LIST_01": "GET /api/books?page=1&limit=10\nKỳ vọng books.length<=10, pagination đúng.",
    "TC_BOOK_LIST_02": "GET /api/books?page=3&limit=20\nKỳ vọng tương đương OFFSET 40 LIMIT 20.",
    "TC_BOOK_LIST_03": "GET /api/books?search=clean\nCheckDB query LIKE title/isbn/publisher/description/author.",
    "TC_BOOK_LIST_04": "GET /api/books?search=clean&availableAt=ebook&availableAt=book-copy\nKỳ vọng vừa match search vừa availability.",
    "TC_BOOK_LIST_05": "GET /api/books?sortBy=publishYear&sortOrder=desc\nCheckDB ORDER BY publishYear DESC.",
    "TC_BOOK_LIST_06": "GET /api/books?authorIds=1&categoryIds=2&languageCodes=vi\nKỳ vọng thỏa 3 filter.",
}

for sheet_name, mapping in [("3.TC_XacThuc", postman_auth), ("4.TC_TimKiemSach", postman_book)]:
    if sheet_name not in wb.sheetnames:
        continue
    ws = wb[sheet_name]
    new_col = ws.max_column + 1
    ws.cell(row=2, column=new_col, value="Lệnh Postman / CheckDB SQL")
    ws.cell(row=2, column=new_col).font = Font(bold=True)
    ws.cell(row=2, column=new_col).fill = header_fill
    ws.cell(row=2, column=new_col).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.cell(row=2, column=new_col).border = thin_border
    ws.column_dimensions[get_column_letter(new_col)].width = 74

    for row in range(3, ws.max_row + 1):
        tc_id = ws.cell(row=row, column=1).value
        if tc_id in mapping:
            c = ws.cell(row=row, column=new_col, value=mapping[tc_id])
            c.alignment = cell_align
            c.border = thin_border
            c.fill = postman_fill

if "8.Postman_CheckDB" in wb.sheetnames:
    del wb["8.Postman_CheckDB"]
ws = wb.create_sheet("8.Postman_CheckDB")
for col, width in {"A": 8, "B": 28, "C": 42, "D": 65, "E": 55, "F": 35}.items():
    ws.column_dimensions[col].width = width

ws.merge_cells("A1:F1")
ws["A1"] = "HƯỚNG DẪN POSTMAN & CHECKDB"
ws["A1"].font = Font(bold=True, size=14, color="FFFFFF")
ws["A1"].fill = title_fill
ws["A1"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

headers = ["STT", "Nhóm", "Mục tiêu", "Postman Request", "CheckDB SQL", "Rollback/Cleanup"]
for i, h in enumerate(headers, 1):
    cell = ws.cell(row=2, column=i, value=h)
    cell.font = Font(bold=True)
    cell.fill = header_fill
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = thin_border

rows = [
    [1, "Auth/Register", "Kiểm tra user đã được chèn", "POST {{baseUrl}}/api/auth/register", "SELECT id,email,status,role FROM User WHERE email='nguyenvana_postman@example.com';", "DELETE token + user test sau khi xong"],
    [2, "Auth/Login", "Kiểm tra RefreshToken đã tạo", "POST {{baseUrl}}/api/auth/login", "SELECT * FROM RefreshToken WHERE userId=(SELECT id FROM User WHERE email='nguyenvana_postman@example.com');", "DELETE FROM RefreshToken WHERE userId=<id>;"],
    [3, "Auth/Logout", "Kiểm tra token đã xóa", "POST {{baseUrl}}/api/auth/logout", "SELECT COUNT(*) FROM RefreshToken WHERE token='<refreshToken>';", "Không rollback (xóa token là đúng hành vi)"],
    [4, "Auth/ChangePassword", "Kiểm tra hash đổi + token clear", "POST {{baseUrl}}/api/auth/change-password", "SELECT password FROM User WHERE email='nguyenvana_postman@example.com'; SELECT COUNT(*) FROM RefreshToken WHERE userId=<id>;", "Đổi lại mật khẩu cũ nếu cần"],
    [5, "Books/List", "Kiểm tra phân trang", "GET {{baseUrl}}/api/books?page=1&limit=10", "SELECT COUNT(*) FROM Book WHERE isDeleted=0;", "Không rollback (chỉ đọc)"],
    [6, "Books/Search", "Kiểm tra keyword search", "GET {{baseUrl}}/api/books?search=clean", "SELECT ... WHERE title/isbn/publisher/description/author LIKE '%clean%';", "Không rollback"],
    [7, "Books/Filter", "Kiểm tra ebook filter", "GET {{baseUrl}}/api/books?availableAt=ebook", "SELECT DISTINCT bookId FROM BookEdition WHERE format='EBOOK' AND isDeleted=0;", "Không rollback"],
    [8, "Books/Sort", "Kiểm tra sort publishYear", "GET {{baseUrl}}/api/books?sortBy=publishYear&sortOrder=desc", "SELECT publishYear FROM Book WHERE isDeleted=0 ORDER BY publishYear DESC;", "Không rollback"],
]
for r_idx, row in enumerate(rows, 3):
    for c_idx, val in enumerate(row, 1):
        cell = ws.cell(row=r_idx, column=c_idx, value=val)
        cell.alignment = center_align if c_idx == 1 else cell_align
        cell.border = thin_border
        if c_idx >= 4:
            cell.fill = postman_fill

start = len(rows) + 5
ws.merge_cells(start_row=start, start_column=1, end_row=start, end_column=6)
ws.cell(row=start, column=1, value="Biến Postman gợi ý").font = Font(bold=True)
ws.cell(row=start, column=1).fill = section_fill
for c in range(1, 7):
    ws.cell(row=start, column=c).border = thin_border

vars_rows = [
    ["baseUrl", base_url, "URL local server"],
    ["accessToken", "<token từ login>", "Dùng auth endpoint"],
    ["refreshToken", "<token từ login>", "Dùng refresh/logout"],
    ["testEmail", "nguyenvana_postman@example.com", "Email riêng để cleanup"],
]
for i, row in enumerate(vars_rows, start + 1):
    ws.cell(row=i, column=1, value=row[0]).border = thin_border
    ws.cell(row=i, column=2, value=row[1]).border = thin_border
    ws.cell(row=i, column=3, value=row[2]).border = thin_border
    for c in range(1, 4):
        ws.cell(row=i, column=c).alignment = cell_align

wb.save(out)
print(f"DONE: {out}")