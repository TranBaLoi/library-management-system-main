import sys
sys.stdout.reconfigure(encoding="utf-8")
from openpyxl import load_workbook
p = r"C:\Users\Loi\Desktop\SQA\10_unit_test.xlsx"
wb = load_workbook(p, data_only=True)
for sn in ["3.TC_XacThuc", "4.TC_TimKiemSach"]:
    if sn in wb.sheetnames:
        ws = wb[sn]
        print(f"\n--- {sn} ---")
        for r in range(2, min(ws.max_row, 30) + 1):
            id_val = ws.cell(r, 1).value
            title_val = ws.cell(r, 2).value
            if id_val:
                print(f"{id_val} | {title_val}")