import sys
sys.stdout.reconfigure(encoding="utf-8")
from openpyxl import load_workbook
p = r"C:\Users\Loi\Desktop\SQA\10_unit_test.xlsx"
wb = load_workbook(p, data_only=True)
print("Sheets:", wb.sheetnames)
for sn in wb.sheetnames:
    ws = wb[sn]
    ids = []
    for r in range(1, min(ws.max_row, 300)+1):
        v = ws.cell(r,1).value
        if isinstance(v, str) and v.strip().startswith(("TC", "UT")):
            ids.append(v.strip())
    if ids:
        print(f"\nSheet {sn} (Col A) IDs ({len(ids)}):")
        print(ids)