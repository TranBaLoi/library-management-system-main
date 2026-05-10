import sys
sys.stdout.reconfigure(encoding="utf-8")
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
import os

wb = Workbook()

# Styles
thin_border = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin")
)
title_font = Font(bold=True, size=14, color="FFFFFF")
title_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
title_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
header_font = Font(bold=True, size=11)
header_fill = PatternFill(start_color="D6E4F0", end_color="D6E4F0", fill_type="solid")
header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
cell_align = Alignment(vertical="top", wrap_text=True)
center_align = Alignment(horizontal="center", vertical="top", wrap_text=True)
pass_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
fail_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
section_font = Font(bold=True, size=12)
section_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
high_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
med_fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
low_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")

def apply_border(ws, row_start, row_end, col_start, col_end):
    for r in range(row_start, row_end + 1):
        for c in range(col_start, col_end + 1):
            cell = ws.cell(row=r, column=c)
            cell.border = thin_border

def set_title(ws, row, col_start, col_end, text):
    ws.merge_cells(start_row=row, start_column=col_start, end_row=row, end_column=col_end)
    cell = ws.cell(row=row, column=col_start)
    cell.value = text
    cell.font = title_font
    cell.fill = title_fill
    cell.alignment = title_align
    for c in range(col_start, col_end + 1):
        ws.cell(row=row, column=c).border = thin_border

def set_header(ws, row, headers):
    for i, h in enumerate(headers, 1):
        cell = ws.cell(row=row, column=i)
        cell.value = h
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border

def set_section(ws, row, col_start, col_end, text):
    ws.merge_cells(start_row=row, start_column=col_start, end_row=row, end_column=col_end)
    cell = ws.cell(row=row, column=col_start)
    cell.value = text
    cell.font = section_font
    cell.fill = section_fill
    cell.alignment = Alignment(vertical="center", wrap_text=True)
    for c in range(col_start, col_end + 1):
        ws.cell(row=row, column=c).border = thin_border

def write_row(ws, row, values, center_cols=set()):
    for i, v in enumerate(values, 1):
        cell = ws.cell(row=row, column=i)
        cell.value = v
        cell.alignment = center_align if i in center_cols else cell_align
        cell.border = thin_border

# =====================================================
# SHEET 1: T\u1ed5ng quan
# =====================================================
ws1 = wb.active
ws1.title = "1.Tong quan"
ws1.column_dimensions["A"].width = 30
ws1.column_dimensions["B"].width = 75
set_title(ws1, 1, 1, 2, "B\u00c1O C\u00c1O KI\u1ec2M TH\u1eec \u0110\u01a0N V\u1eca (UNIT TESTING REPORT)")
info = [
    ("D\u1ef1 \u00e1n", "Library Management System"),
    ("Ng\u00f4n ng\u1eef / Framework", "TypeScript / Next.js 15 + Prisma ORM"),
    ("Testing Framework", "Jest 30 + ts-jest 29"),
    ("C\u00f4ng c\u1ee5 Mock", "jest.fn(), jest.mock() - mock Prisma, bcryptjs, jsonwebtoken"),
    ("C\u00f4ng c\u1ee5 Coverage", "Jest built-in coverage (Istanbul/V8)"),
    ("Ng\u00e0y th\u1ef1c hi\u1ec7n", "08/05/2026"),
    ("Ph\u1ea1m vi ki\u1ec3m th\u1eed", "1) X\u00e1c th\u1ef1c t\u00e0i kho\u1ea3n (AuthService)\n2) T\u00ecm ki\u1ebfm v\u00e0 duy\u1ec7t s\u00e1ch (BookService)"),
]
for i, (k, v) in enumerate(info, 3):
    ws1.cell(row=i, column=1, value=k).font = Font(bold=True, size=11)
    ws1.cell(row=i, column=1).fill = section_fill
    ws1.cell(row=i, column=1).border = thin_border
    ws1.cell(row=i, column=1).alignment = cell_align
    ws1.cell(row=i, column=2, value=v).alignment = cell_align
    ws1.cell(row=i, column=2).border = thin_border

# =====================================================
# SHEET 2: Ph\u1ea1m vi ki\u1ec3m th\u1eed
# =====================================================
ws2 = wb.create_sheet("2.Pham vi")
widths2 = [8, 38, 42, 15, 55]
for i, w in enumerate(widths2, 1):
    ws2.column_dimensions[get_column_letter(i)].width = w

set_title(ws2, 1, 1, 5, "PH\u1ea0M VI KI\u1ec2M TH\u1eec")
set_section(ws2, 3, 1, 5, "2a. C\u00e1c h\u00e0m/class \u0110\u01af\u1ee2C ki\u1ec3m th\u1eed")
set_header(ws2, 4, ["STT", "T\u1ec7p", "H\u00e0m/Ph\u01b0\u01a1ng th\u1ee9c", "Lo\u1ea1i", "M\u00f4 t\u1ea3"])

tested_data = [
    (1,"src/services/auth.service.ts","AuthService.register()","Service","\u0110\u0103ng k\u00fd t\u00e0i kho\u1ea3n m\u1edbi"),
    (2,"src/services/auth.service.ts","AuthService.login()","Service","\u0110\u0103ng nh\u1eadp - x\u00e1c th\u1ef1c email/password, t\u1ea1o JWT"),
    (3,"src/services/auth.service.ts","AuthService.refreshAccessToken()","Service","L\u00e0m m\u1edbi Access Token t\u1eeb Refresh Token"),
    (4,"src/services/auth.service.ts","AuthService.logout()","Service","\u0110\u0103ng xu\u1ea5t - x\u00f3a refresh token kh\u1ecfi DB"),
    (5,"src/services/auth.service.ts","AuthService.logoutAll()","Service","\u0110\u0103ng xu\u1ea5t t\u1ea5t c\u1ea3 thi\u1ebft b\u1ecb"),
    (6,"src/services/auth.service.ts","AuthService.changePassword()","Service","\u0110\u1ed5i m\u1eadt kh\u1ea9u - c\u1eadp nh\u1eadt hash, logout all"),
    (7,"src/services/auth.service.ts","AuthService.cleanupExpiredTokens()","Service","D\u1ecdn d\u1eb9p token h\u1ebft h\u1ea1n"),
    (8,"src/services/book.service.ts","transformBookData()","Utility","Chu\u1ea9n h\u00f3a d\u1eef li\u1ec7u s\u00e1ch t\u1eeb Prisma"),
    (9,"src/services/book.service.ts","buildBookWhereClause()","Utility","T\u1ea1o \u0111i\u1ec1u ki\u1ec7n WHERE t\u1eeb filter params"),
    (10,"src/services/book.service.ts","buildOrderByClause()","Utility","T\u1ea1o ORDER BY t\u1eeb sort params"),
    (11,"src/services/book.service.ts","listBooks()","Service","T\u00ecm ki\u1ebfm & ph\u00e2n trang s\u00e1ch"),
]
for i, row in enumerate(tested_data):
    write_row(ws2, 5 + i, row, center_cols={1, 4})

nr = 5 + len(tested_data) + 1
set_section(ws2, nr, 1, 5, "2b. C\u00e1c h\u00e0m/class KH\u00d4NG c\u1ea7n ki\u1ec3m th\u1eed")
set_header(ws2, nr+1, ["STT", "T\u1ec7p", "H\u00e0m/Module", "Lo\u1ea1i", "L\u00fd do kh\u00f4ng ki\u1ec3m th\u1eed"])
not_tested_data = [
    (1,"src/components/","React Components","UI","Thu\u1ed9c scope Integration/E2E test"),
    (2,"src/lib/hooks/","Custom React Hooks","UI Hook","Ph\u1ee5 thu\u1ed9c React context/DOM"),
    (3,"src/app/api/*/route.ts","API Route handlers","Route","L\u1edbp HTTP, thu\u1ed9c integration test"),
    (4,"src/middleware.ts","Next.js Middleware","Middleware","Ph\u1ee5 thu\u1ed9c Edge runtime"),
    (5,"src/services/gorse.service.ts","GorseService","Service","G\u1ecdi API b\u00ean ngo\u00e0i"),
    (6,"src/services/qdrant.service.ts","QdrantService","Service","G\u1ecdi API vector DB b\u00ean ngo\u00e0i"),
    (7,"src/workers/","Background Workers","Worker","C\u1ea7n integration test v\u1edbi Redis"),
    (8,"src/lib/cron/","Cron Jobs","Cron","Scheduled tasks"),
]
for i, row in enumerate(not_tested_data):
    write_row(ws2, nr+2+i, row, center_cols={1, 4})

print("Sheet 1-2 done")

# =====================================================
# SHEET 3: TC Auth
# =====================================================
ws3 = wb.create_sheet("3.TC_XacThuc")
widths3 = [22, 45, 50, 45, 12, 40]
for i, w in enumerate(widths3, 1):
    ws3.column_dimensions[get_column_letter(i)].width = w

set_title(ws3, 1, 1, 6, "TEST CASES - X\u00c1C TH\u1ef0C T\u00c0I KHO\u1ea2N")
set_header(ws3, 2, ["Test Case ID", "M\u1ee5c ti\u00eau", "Input", "Expected Output", "K\u1ebft qu\u1ea3", "Ghi ch\u00fa"])

auth_cases = [
    ("TC_AUTH_REG_01","\u0110\u0103ng k\u00fd th\u00e0nh c\u00f4ng v\u1edbi d\u1eef li\u1ec7u h\u1ee3p l\u1ec7","fullName='Nguyen Van A', email='nguyenvana@example.com', password='Password123!', confirmPassword='Password123!'","Tr\u1ea3 v\u1ec1 {user, message}, user.role=READER, password \u0111\u01b0\u1ee3c hash","PASS","CheckDB: create g\u1ecdi 1 l\u1ea7n, pw kh\u00f4ng plaintext"),
    ("TC_AUTH_REG_02","Email sai \u0111\u1ecbnh d\u1ea1ng","email='invalid-email'","N\u00e9m ValidationError 'Invalid email'","PASS","create KH\u00d4NG g\u1ecdi"),
    ("TC_AUTH_REG_03","M\u1eadt kh\u1ea9u kh\u00f4ng kh\u1edbp confirm","password='Password123!' vs confirmPassword='Different456@'","N\u00e9m ValidationError 'do not match'","PASS","Validate tr\u01b0\u1edbc DB"),
    ("TC_AUTH_REG_04","Email \u0111\u00e3 t\u1ed3n t\u1ea1i","email \u0111\u00e3 c\u00f3 trong DB","N\u00e9m ConflictError 'already registered'","PASS","CheckDB: findUnique\u2192create KH\u00d4NG g\u1ecdi"),
    ("TC_AUTH_REG_05","M\u1eadt kh\u1ea9u y\u1ebfu","password='12345678'","N\u00e9m ValidationError (thi\u1ebfu ch\u1eef hoa/k\u00fd t\u1ef1 \u0111\u1eb7c bi\u1ec7t)","PASS","Ki\u1ec3m tra quy t\u1eafc password"),
    ("TC_AUTH_REG_06","H\u1ecd t\u00ean r\u1ed7ng","fullName=''","N\u00e9m ValidationError","PASS","B\u1eaft bu\u1ed9c fullName"),
    ("TC_AUTH_REG_07","\u0110\u0103ng k\u00fd k\u00e8m s\u0111t h\u1ee3p l\u1ec7","phoneNumber='0901234567'","Th\u00e0nh c\u00f4ng, user.phoneNumber='0901234567'","PASS","Optional field"),
    ("TC_AUTH_REG_08","S\u0110T kh\u00f4ng \u0111\u1ee7 10 ch\u1eef s\u1ed1","phoneNumber='123'","N\u00e9m ValidationError","PASS","Validate phone"),
    ("TC_AUTH_LOGIN_01","\u0110\u0103ng nh\u1eadp th\u00e0nh c\u00f4ng","email/password \u0111\u00fang, user ACTIVE","Tr\u1ea3 {userId, accessToken, refreshToken}, refreshToken l\u01b0u DB","PASS","CheckDB: refreshToken.create + user.update"),
    ("TC_AUTH_LOGIN_02","Email kh\u00f4ng t\u1ed3n t\u1ea1i","email='notexist@example.com'","N\u00e9m UnauthorizedError","PASS","Kh\u00f4ng ti\u1ebft l\u1ed9 info"),
    ("TC_AUTH_LOGIN_03","M\u1eadt kh\u1ea9u sai","email \u0111\u00fang, password sai","N\u00e9m UnauthorizedError","PASS","Kh\u00f4ng leak sai password"),
    ("TC_AUTH_LOGIN_04","T\u00e0i kho\u1ea3n INACTIVE","user.status=INACTIVE","N\u00e9m UnauthorizedError 'inactive'","PASS","Ch\u1eb7n inactive"),
    ("TC_AUTH_LOGIN_05","T\u00e0i kho\u1ea3n \u0111\u00e3 x\u00f3a m\u1ec1m","user.isDeleted=true","N\u00e9m UnauthorizedError","PASS","Ch\u1eb7n deleted"),
    ("TC_AUTH_LOGIN_06","Email login sai format","email='not_an_email'","N\u00e9m ValidationError","PASS","Validate \u0111\u1ea7u v\u00e0o"),
    ("TC_AUTH_LOGIN_07","Password r\u1ed7ng","password=''","N\u00e9m ValidationError","PASS","Validate \u0111\u1ea7u v\u00e0o"),
    ("TC_AUTH_LOGIN_08","First login","firstLoginAt=null","isFirstLogin=true, update firstLoginAt","PASS","CheckDB: user.update"),
    ("TC_AUTH_LOGIN_09","Not first login","firstLoginAt c\u00f3 gi\u00e1 tr\u1ecb","isFirstLogin=false, KH\u00d4NG update","PASS","Gi\u1eef nguy\u00ean firstLoginAt"),
    ("TC_AUTH_LOGIN_10","rememberMe=true","rememberMe=true","expiresAt ~30 ng\u00e0y","PASS","CheckDB: expiresAt trong create"),
    ("TC_AUTH_REFRESH_01","L\u00e0m m\u1edbi token th\u00e0nh c\u00f4ng","refreshToken h\u1ee3p l\u1ec7, user ACTIVE","Tr\u1ea3 accessToken m\u1edbi","PASS","jwt.sign g\u1ecdi"),
    ("TC_AUTH_REFRESH_02","Token kh\u00f4ng t\u1ed3n t\u1ea1i DB","tokenId kh\u00f4ng c\u00f3 DB","N\u00e9m UnauthorizedError","PASS","findUnique null"),
    ("TC_AUTH_REFRESH_03","Token h\u1ebft h\u1ea1n","expiresAt < now","N\u00e9m UnauthorizedError, x\u00f3a token DB","PASS","CheckDB: delete g\u1ecdi"),
    ("TC_AUTH_REFRESH_04","User inactive","user.status=INACTIVE","N\u00e9m UnauthorizedError","PASS","Ch\u1eb7n inactive"),
    ("TC_AUTH_CHANGEPWD_01","\u0110\u1ed5i m\u1eadt kh\u1ea9u th\u00e0nh c\u00f4ng","current \u0111\u00fang, new h\u1ee3p l\u1ec7","Update password, logoutAll","PASS","CheckDB: update + deleteMany"),
    ("TC_AUTH_CHANGEPWD_02","Current password sai","currentPassword sai","N\u00e9m ValidationError","PASS","Kh\u00f4ng update"),
    ("TC_AUTH_CHANGEPWD_03","Confirm kh\u00f4ng kh\u1edbp","newPassword != confirm","N\u00e9m ValidationError","PASS","Kh\u00f4ng update"),
    ("TC_AUTH_CHANGEPWD_04","M\u1eadt kh\u1ea9u m\u1edbi tr\u00f9ng c\u0169","new == old","N\u00e9m ValidationError 'different'","PASS","Kh\u00f4ng update"),
    ("TC_AUTH_CHANGEPWD_05","User kh\u00f4ng t\u1ed3n t\u1ea1i","userId=9999","N\u00e9m NotFoundError","PASS","findUnique null"),
    ("TC_AUTH_LOGOUT_01","Logout th\u00e0nh c\u00f4ng","refreshToken h\u1ee3p l\u1ec7","X\u00f3a token kh\u1ecfi DB","PASS","Rollback: delete token"),
    ("TC_AUTH_LOGOUT_02","Logout t\u1ea5t c\u1ea3 thi\u1ebft b\u1ecb","userId=1","X\u00f3a m\u1ecdi token c\u1ee7a user","PASS","Rollback: deleteMany"),
    ("TC_AUTH_CLEANUP_01","D\u1ecdn token h\u1ebft h\u1ea1n","DB c\u00f3 5 token expired","Tr\u1ea3 v\u1ec1 5","PASS","deleteMany expiresAt<now"),
    ("TC_AUTH_CLEANUP_02","Kh\u00f4ng c\u00f3 token h\u1ebft h\u1ea1n","DB r\u1ed7ng","Tr\u1ea3 v\u1ec1 0","PASS","No-op"),
]
for i, row in enumerate(auth_cases):
    r = 3 + i
    write_row(ws3, r, row, center_cols={1, 5})
    c5 = ws3.cell(row=r, column=5)
    c5.fill = pass_fill if row[4] == "PASS" else fail_fill

print("Sheet 3 done")

# =====================================================
# SHEET 4: TC Book
# =====================================================
ws4 = wb.create_sheet("4.TC_TimKiemSach")
for i, w in enumerate(widths3, 1):
    ws4.column_dimensions[get_column_letter(i)].width = w

set_title(ws4, 1, 1, 6, "TEST CASES - T\u00ccM KI\u1ebeM V\u00c0 DUY\u1ec6T S\u00c1CH")
set_header(ws4, 2, ["Test Case ID", "M\u1ee5c ti\u00eau", "Input", "Expected Output", "K\u1ebft qu\u1ea3", "Ghi ch\u00fa"])

book_cases = [
    ("TC_BOOK_TRANSFORM_01","Transform d\u1eef li\u1ec7u \u0111\u1ea7y \u0111\u1ee7","BookRawData: 2 EBOOK, 1 AUDIO, 3 reviews(5,4,5), 5 items, 2 categories","categories/count/rating \u0111\u00fang, averageRating=4.7","PASS","rating l\u00e0m tr\u00f2n 1 ch\u1eef s\u1ed1"),
    ("TC_BOOK_TRANSFORM_02","Kh\u00f4ng c\u00f3 review","reviews=[]","averageRating=0","PASS","Tr\u00e1nh chia cho 0"),
    ("TC_BOOK_TRANSFORM_03","Thi\u1ebfu d\u1eef li\u1ec7u optional","bookCategories/bookEditions/count = undefined","categories=[], count=0, ebook=0, audio=0","PASS","Nullish coalescing"),
    ("TC_BOOK_FILTER_01","M\u1eb7c \u0111\u1ecbnh s\u00e1ch ch\u01b0a x\u00f3a","isDeleted=null","where.isDeleted=false","PASS","Soft delete default"),
    ("TC_BOOK_FILTER_02","L\u1ecdc theo t\u00e1c gi\u1ea3","authorIds=[1,2,3]","where.authorId={in:[1,2,3]}","PASS","Multi-author filter"),
    ("TC_BOOK_FILTER_03","L\u1ecdc theo danh m\u1ee5c","categoryIds=[5,6]","where.bookCategories.some.categoryId.in","PASS","Relation filter"),
    ("TC_BOOK_FILTER_04","L\u1ecdc theo ng\u00f4n ng\u1eef","languageCodes=['vi','en']","where.language={in:['vi','en']}","PASS","Multi-language"),
    ("TC_BOOK_FILTER_05","L\u1ecdc n\u0103m xu\u1ea5t b\u1ea3n","from=2000, to=2020","where.publishYear={gte:2000,lte:2020}","PASS","Range filter"),
    ("TC_BOOK_FILTER_06","L\u1ecdc s\u00e1ch c\u00f3 ebook","availableAt=['ebook']","where.bookEditions.some EBOOK","PASS","Edition filter"),
    ("TC_BOOK_FILTER_07","L\u1ecdc s\u00e1ch c\u00f3 b\u1ea3n in","availableAt=['book-copy']","where.bookItems.some AVAILABLE","PASS","Inventory filter"),
    ("TC_BOOK_FILTER_08","C\u1ea3 ebook v\u00e0 b\u1ea3n in","availableAt=['ebook','book-copy']","where.AND ch\u1ee9a 2 \u0111i\u1ec1u ki\u1ec7n","PASS","Combined"),
    ("TC_BOOK_SORT_01","Sort m\u1eb7c \u0111\u1ecbnh","kh\u00f4ng truy\u1ec1n sort","{createdAt:'desc'}","PASS","Default sort"),
    ("TC_BOOK_SORT_02","Sort title asc","sortBy='title', sortOrder='asc'","{title:'asc'}","PASS","Title sort"),
    ("TC_BOOK_SORT_03","Sort price desc","sortBy='price', sortOrder='desc'","{price:'desc'}","PASS","Price sort"),
    ("TC_BOOK_SORT_04","Sort field kh\u00f4ng h\u1ee3p l\u1ec7","sortBy='unknownField'","{createdAt:'desc'}","PASS","Fallback"),
    ("TC_BOOK_LIST_01","Danh s\u00e1ch trang 1","page=1, limit=10","skip=0, take=10","PASS","CheckDB: findMany+count"),
    ("TC_BOOK_LIST_02","Ph\u00e2n trang page=3","page=3, limit=20","skip=40, take=20","PASS","Offset \u0111\u00fang"),
    ("TC_BOOK_LIST_03","T\u00ecm theo t\u1eeb kh\u00f3a","search='clean'","where.OR contains on title/isbn/publisher/description/author","PASS","SQL LIKE"),
    ("TC_BOOK_LIST_04","Search + availability","search='clean' + availableAt both","where.AND with OR","PASS","Complex query"),
    ("TC_BOOK_LIST_05","Sort publishYear desc","sortBy='publishYear', sortOrder='desc'","orderBy publishYear desc","PASS","Sort mapping"),
    ("TC_BOOK_LIST_06","Multi filter","authorIds+categoryIds+languageCodes","where \u0111\u1ee7 3 \u0111i\u1ec1u ki\u1ec7n","PASS","Combined filter"),
]
for i, row in enumerate(book_cases):
    r = 3 + i
    write_row(ws4, r, row, center_cols={1, 5})
    c5 = ws4.cell(row=r, column=5)
    c5.fill = pass_fill if row[4] == "PASS" else fail_fill

print("Sheet 4 done")

# =====================================================
# SHEET 5: Execution Report
# =====================================================
ws5 = wb.create_sheet("5.Ket qua thuc thi")
widths5 = [42, 15, 15, 15, 15]
for i, w in enumerate(widths5, 1):
    ws5.column_dimensions[get_column_letter(i)].width = w

set_title(ws5, 1, 1, 5, "K\u1ebeT QU\u1ea2 TH\u1ef0C THI KI\u1ec2M TH\u1eec")
set_header(ws5, 2, ["Test Suite", "S\u1ed1 Test", "Pass", "Fail", "T\u1ef7 l\u1ec7 Pass"])

exec_rows = [
    ("TC_AUTH_REG (\u0110\u0103ng k\u00fd)", 8, 8, 0, "100%"),
    ("TC_AUTH_LOGIN (\u0110\u0103ng nh\u1eadp)", 10, 10, 0, "100%"),
    ("TC_AUTH_REFRESH (L\u00e0m m\u1edbi token)", 4, 4, 0, "100%"),
    ("TC_AUTH_CHANGEPWD (\u0110\u1ed5i m\u1eadt kh\u1ea9u)", 5, 5, 0, "100%"),
    ("TC_AUTH_LOGOUT (\u0110\u0103ng xu\u1ea5t)", 2, 2, 0, "100%"),
    ("TC_AUTH_CLEANUP (D\u1ecdn token)", 2, 2, 0, "100%"),
    ("TC_BOOK_TRANSFORM (Chu\u1ea9n h\u00f3a data)", 3, 3, 0, "100%"),
    ("TC_BOOK_FILTER (\u0110i\u1ec1u ki\u1ec7n l\u1ecdc)", 8, 8, 0, "100%"),
    ("TC_BOOK_SORT (S\u1eafp x\u1ebfp)", 4, 4, 0, "100%"),
    ("TC_BOOK_LIST (T\u00ecm ki\u1ebfm & ph\u00e2n trang)", 6, 6, 0, "100%"),
]
for i, row in enumerate(exec_rows):
    r = 3 + i
    write_row(ws5, r, row, center_cols={2, 3, 4, 5})
    ws5.cell(row=r, column=5).fill = pass_fill

total_r = 3 + len(exec_rows)
set_section(ws5, total_r, 1, 1, "T\u1ed4NG C\u1ed8NG")
for j, v in enumerate([52, 52, 0, "100%"], 2):
    c = ws5.cell(row=total_r, column=j, value=v)
    c.font = Font(bold=True, size=11)
    c.fill = pass_fill
    c.alignment = center_align
    c.border = thin_border

ws5.cell(row=total_r+2, column=1, value="L\u1ec7nh ch\u1ea1y test:").font = Font(bold=True)
ws5.cell(row=total_r+2, column=2, value="npx jest --config jest.config.ts --coverage --verbose")
ws5.cell(row=total_r+3, column=1, value="K\u1ebft qu\u1ea3:").font = Font(bold=True)
ws5.cell(row=total_r+3, column=2, value="Test Suites: 2 passed, 2 total | Tests: 52 passed, 52 total | Time: 2.344s")
ws5.cell(row=total_r+5, column=1, value="(\u1ea2nh ch\u1ee5p m\u00e0n h\u00ecnh k\u1ebft qu\u1ea3 test \u0111\u01b0\u1ee3c \u0111\u00ednh k\u00e8m trong th\u01b0 m\u1ee5c coverage/)")

print("Sheet 5 done")

# =====================================================
# SHEET 6: Coverage
# =====================================================
ws6 = wb.create_sheet("6.Code Coverage")
widths6 = [38, 16, 16, 16, 16, 40]
for i, w in enumerate(widths6, 1):
    ws6.column_dimensions[get_column_letter(i)].width = w

set_title(ws6, 1, 1, 6, "B\u00c1O C\u00c1O \u0110\u1ed8 BAO PH\u1ee6 M\u00c3 NGU\u1ed2N")
set_header(ws6, 2, ["T\u1ec7p", "% Statements", "% Branches", "% Functions", "% Lines", "Ghi ch\u00fa"])

cov_data = [
    ("services/auth.service.ts", "97.82%", "95.08%", "100%", "97.82%", "Uncovered: line 108 (catch GorseService log)"),
    ("services/book.service.ts", "93.84%", "90.47%", "100%", "93.65%", "Uncovered: line 113-118 (edge availability)"),
    ("lib/errors/errors.ts", "82.60%", "0%", "71.42%", "82.60%", "Ch\u1ec9 2 error class kh\u00f4ng \u0111\u01b0\u1ee3c t\u1ea1o"),
    ("types/book.ts", "100%", "100%", "100%", "100%", ""),
]

def cov_fill(pct_str):
    try:
        p = float(pct_str.replace("%",""))
        if p >= 90: return high_fill
        if p >= 70: return med_fill
        return low_fill
    except: return None

for i, row in enumerate(cov_data):
    r = 3 + i
    write_row(ws6, r, row, center_cols={2,3,4,5})
    for j in range(2, 6):
        f = cov_fill(str(row[j-1]))
        if f: ws6.cell(row=r, column=j).fill = f

sr = 3 + len(cov_data) + 1
set_section(ws6, sr, 1, 6, "T\u00d3M T\u1eaeT: Coverage cho 2 t\u00ednh n\u0103ng ch\u00ednh")
ws6.cell(row=sr+1, column=1, value="auth.service.ts").font = Font(bold=True)
ws6.cell(row=sr+1, column=2, value="97.82% Stmts").fill = high_fill
ws6.cell(row=sr+1, column=3, value="95.08% Branch").fill = high_fill
ws6.cell(row=sr+1, column=4, value="100% Funcs").fill = high_fill
ws6.cell(row=sr+1, column=5, value="97.82% Lines").fill = high_fill
ws6.cell(row=sr+2, column=1, value="book.service.ts").font = Font(bold=True)
ws6.cell(row=sr+2, column=2, value="93.84% Stmts").fill = high_fill
ws6.cell(row=sr+2, column=3, value="90.47% Branch").fill = high_fill
ws6.cell(row=sr+2, column=4, value="100% Funcs").fill = high_fill
ws6.cell(row=sr+2, column=5, value="93.65% Lines").fill = high_fill

ws6.cell(row=sr+4, column=1, value="C\u00f4ng c\u1ee5:").font = Font(bold=True)
ws6.cell(row=sr+4, column=2, value="Jest built-in coverage (Istanbul/V8)")
ws6.cell(row=sr+5, column=1, value="Report HTML:").font = Font(bold=True)
ws6.cell(row=sr+5, column=2, value="coverage/lcov-report/index.html")

print("Sheet 6 done")

# =====================================================
# SHEET 7: T\u00e0i li\u1ec7u & Prompt
# =====================================================
ws7 = wb.create_sheet("7.Tai lieu Prompt")
ws7.column_dimensions["A"].width = 8
ws7.column_dimensions["B"].width = 85

set_title(ws7, 1, 1, 2, "T\u00c0I LI\u1ec6U THAM KH\u1ea2O & DANH S\u00c1CH PROMPT")
set_section(ws7, 3, 1, 2, "7a. T\u00e0i li\u1ec7u tham kh\u1ea3o")
refs = [
    "Jest Documentation - https://jestjs.io/docs/getting-started",
    "ts-jest - TypeScript preprocessor - https://kulshekhar.github.io/ts-jest/",
    "Prisma Unit Testing - https://www.prisma.io/docs/guides/testing/unit-testing",
    "Jest Mock Functions - https://jestjs.io/docs/mock-functions",
    "bcryptjs - https://www.npmjs.com/package/bcryptjs",
    "jsonwebtoken - https://www.npmjs.com/package/jsonwebtoken",
    "Next.js 15 App Router - https://nextjs.org/docs/app/building-your-application",
    "T\u00e0i li\u1ec7u m\u00f4n \u0110\u1ea3m b\u1ea3o ch\u1ea5t l\u01b0\u1ee3ng ph\u1ea7n m\u1ec1m (SQA)",
]
for i, ref in enumerate(refs):
    write_row(ws7, 4 + i, [i + 1, ref], center_cols={1})

pr = 4 + len(refs) + 1
set_section(ws7, pr, 1, 2, "7b. Danh s\u00e1ch prompt \u0111\u00e3 s\u1eed d\u1ee5ng")
prompts = [
    "T\u1ea1o cho t\u00f4i 1 file excel v\u1edbi y\u00eau c\u1ea7u ki\u1ec3m th\u1eed unit test trong d\u1ef1 \u00e1n library... (prompt g\u1ed1c \u0111\u1ea7y \u0111\u1ee7)",
    "Y\u00eau c\u1ea7u quan tr\u1ecdng: vi\u1ebft b\u1eb1ng ti\u1ebfng Vi\u1ec7t c\u00f3 d\u1ea5u",
    "\u0110\u1ecdc to\u00e0n b\u1ed9 src code v\u00e0 t\u1ea1o test case chi ti\u1ebft cho h\u00e0m/class",
    "T\u1ea1o script test theo m\u00e3 test case + \u0111o coverage",
    "Th\u1ef1c hi\u1ec7n cho 2 t\u00ednh n\u0103ng: X\u00e1c th\u1ef1c t\u00e0i kho\u1ea3n, T\u00ecm ki\u1ebfm v\u00e0 duy\u1ec7t s\u00e1ch",
]
for i, p in enumerate(prompts):
    write_row(ws7, pr + 1 + i, [i + 1, p], center_cols={1})

out = r"C:\Users\Loi\Desktop\SQA\UnitTestReport_LibraryMS.xlsx"
wb.save(out)
print(f"\nDONE: {out}")