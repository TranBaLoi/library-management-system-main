# -*- coding: utf-8 -*-
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

src = r"C:\Users\Loi\Desktop\SQA\UnitTestReport_LibraryMS_v2.xlsx"
out = r"C:\Users\Loi\Desktop\SQA\UnitTestReport_LibraryMS_Postman.xlsx"
wb = load_workbook(src)

thin_border = Border(left=Side(style="thin"), right=Side(style="thin"), top=Side(style="thin"), bottom=Side(style="thin"))
header_fill = PatternFill(start_color="D6E4F0", end_color="D6E4F0", fill_type="solid")
postman_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
section_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
title_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
cell_align = Alignment(vertical="top", wrap_text=True)
center_align = Alignment(horizontal="center", vertical="top", wrap_text=True)

# Remove the previous Postman guide sheet if it exists
if "8.Postman_CheckDB" in wb.sheetnames:
    del wb["8.Postman_CheckDB"]

ws = wb.create_sheet("8.Postman_CheckDB")

for col, width in {"A": 8, "B": 28, "C": 42, "D": 50, "E": 60, "F": 45}.items():
    ws.column_dimensions[col].width = width

ws.merge_cells("A1:F1")
ws["A1"] = "HƯỚNG DẪN POSTMAN & CHECKDB (DOCKER)"
ws["A1"].font = Font(bold=True, size=14, color="FFFFFF")
ws["A1"].fill = title_fill
ws["A1"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

headers = ["STT", "Nhóm", "Mục tiêu", "Postman Request", "CheckDB SQL (chạy trong DB container)", "Rollback/Cleanup"]
for i, h in enumerate(headers, 1):
    cell = ws.cell(row=2, column=i, value=h)
    cell.font = Font(bold=True)
    cell.fill = header_fill
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = thin_border

docker_exec = "docker exec -it library-mysql mysql -u root -pyour_mysql_root_password library_management -e"

rows = [
    [1, "Auth/Register", "Kiểm tra user đã được chèn", "POST {{baseUrl}}/api/auth/register", f"{docker_exec} \"SELECT id,email,status,role FROM User WHERE email='nguyenvana_postman@example.com';\"", f"{docker_exec} \"DELETE FROM User WHERE email='nguyenvana_postman@example.com';\""],
    [2, "Auth/Login", "Kiểm tra RefreshToken đã tạo", "POST {{baseUrl}}/api/auth/login", f"{docker_exec} \"SELECT * FROM RefreshToken WHERE userId=(SELECT id FROM User WHERE email='nguyenvana_postman@example.com');\"", f"{docker_exec} \"DELETE FROM RefreshToken WHERE userId=<id>;\""],
    [3, "Auth/Logout", "Kiểm tra token đã xóa", "POST {{baseUrl}}/api/auth/logout", f"{docker_exec} \"SELECT COUNT(*) FROM RefreshToken WHERE token='<refreshToken>';\"", "Không rollback (xóa token là đúng hành vi)"],
    [4, "Auth/ChangePassword", "Kiểm tra hash đổi + token clear", "POST {{baseUrl}}/api/auth/change-password", f"{docker_exec} \"SELECT password FROM User WHERE email='nguyenvana_postman@example.com'; SELECT COUNT(*) FROM RefreshToken WHERE userId=<id>;\"", "Đổi lại mật khẩu cũ nếu cần"],
    [5, "Books/List", "Kiểm tra phân trang", "GET {{baseUrl}}/api/books?page=1&limit=10", f"{docker_exec} \"SELECT COUNT(*) FROM Book WHERE isDeleted=0;\"", "Không rollback (chỉ đọc)"],
    [6, "Books/Search", "Kiểm tra keyword search", "GET {{baseUrl}}/api/books?search=clean", f"{docker_exec} \"SELECT * FROM Book b LEFT JOIN Author a ON b.authorId=a.id WHERE b.title LIKE '%clean%' OR b.isbn LIKE '%clean%' OR b.publisher LIKE '%clean%' OR b.description LIKE '%clean%' OR a.fullName LIKE '%clean%';\"", "Không rollback"],
    [7, "Books/Filter", "Kiểm tra ebook filter", "GET {{baseUrl}}/api/books?availableAt=ebook", f"{docker_exec} \"SELECT DISTINCT bookId FROM BookEdition WHERE format='EBOOK' AND isDeleted=0;\"", "Không rollback"],
    [8, "Books/Sort", "Kiểm tra sort publishYear", "GET {{baseUrl}}/api/books?sortBy=publishYear&sortOrder=desc", f"{docker_exec} \"SELECT publishYear FROM Book WHERE isDeleted=0 ORDER BY publishYear DESC;\"", "Không rollback"],
]

for r_idx, row in enumerate(rows, 3):
    for c_idx, val in enumerate(row, 1):
        cell = ws.cell(row=r_idx, column=c_idx, value=val)
        cell.alignment = center_align if c_idx == 1 else cell_align
        cell.border = thin_border
        if c_idx >= 4:
            cell.fill = postman_fill

start = len(rows) + 5
ws.merge_cells(start_row=start, start_column=1, end_row=start, end_column=6)
ws.cell(row=start, column=1, value="Hướng dẫn truy cập DB qua Docker").font = Font(bold=True)
ws.cell(row=start, column=1).fill = section_fill
for c in range(1, 7):
    ws.cell(row=start, column=c).border = thin_border

vars_rows = [
    ["Mở MySQL Interactive Shell", "docker exec -it library-mysql mysql -u root -pyour_mysql_root_password library_management", ""],
    ["Xem cấu trúc bảng User", "DESCRIBE User;", ""],
    ["Xem các Refresh Token", "SELECT * FROM RefreshToken LIMIT 5;", ""],
]
for i, row in enumerate(vars_rows, start + 1):
    ws.cell(row=i, column=1, value=row[0]).border = thin_border
    ws.merge_cells(start_row=i, start_column=2, end_row=i, end_column=6)
    ws.cell(row=i, column=2, value=row[1]).border = thin_border
    ws.cell(row=i, column=2).alignment = cell_align
    for c in range(2, 7):
        ws.cell(row=i, column=c).border = thin_border

wb.save(out)
print(f"DONE: {out}")