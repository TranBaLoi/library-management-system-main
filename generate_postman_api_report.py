import os, re, json, subprocess, datetime
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter

ROOT = Path(r'C:\Users\Loi\Desktop\SQA\library-management-system-main')
API_ROOT = ROOT / 'src' / 'app' / 'api'
OUT = ROOT / 'Postman_API_Test_Report_LibraryMS.xlsx'
METHODS = ['GET','POST','PUT','PATCH','DELETE']

body_samples = {
    'auth/register': {'fullName':'Nguyen Van Test','email':'postman_test_{{$timestamp}}@example.com','password':'Password@123','confirmPassword':'Password@123'},
    'auth/login': {'email':'{{reader_email}}','password':'{{reader_password}}'},
    'auth/change-password': {'currentPassword':'{{current_password}}','newPassword':'NewPassword@123','confirmPassword':'NewPassword@123'},
    'auth/reset-password': {'email':'{{reader_email}}','otp':'{{otp_code}}','newPassword':'NewPassword@123','confirmPassword':'NewPassword@123'},
    'auth/interests': {'categoryIds':[1,2,3]},
    'otp/send': {'email':'{{reader_email}}','purpose':'forgot-password'},
    'otp/verify/forgot-password': {'email':'{{reader_email}}','otp':'{{otp_code}}'},
    'authors': {'fullName':'Tac gia API Test {{$timestamp}}','bio':'Created by Postman API test','birthDate':'1980-01-01','nationality':'Vietnam','isDeleted':False},
    'categories': {'name':'Danh muc API Test {{$timestamp}}','description':'Created by Postman API test','isDeleted':False},
    'books': {'authorId':1,'title':'Sach API Test {{$timestamp}}','isbn':'ISBN-{{$timestamp}}','publishYear':2024,'publisher':'NXB Test','pageCount':250,'price':100000,'edition':'1st','description':'Created by Postman API test','language':'vi','categories':[1]},
    'book-editions': {'bookId':1,'format':'EBOOK','fileUrl':'/uploads/test.pdf','fileSize':1024,'isDeleted':False},
    'book-items': {'bookId':1,'code':'COPY-{{$timestamp}}','condition':'GOOD','status':'AVAILABLE','acquisitionDate':'2024-01-01','isDeleted':False},
    'borrow-requests': {'bookId':1,'borrowDate':'2026-05-09','dueDate':'2026-05-23','note':'API test borrow request'},
    'borrow-requests/[id]/manage': {'status':'APPROVED','note':'Approved by API test'},
    'borrow-records/[id]/renew': {'newDueDate':'2026-06-01','note':'Renew by API test'},
    'borrow-records/[id]/return': {'condition':'GOOD','note':'Return by API test'},
    'borrow-records/[id]/return-ebook': {'note':'Return ebook by API test'},
    'ebook-borrow-requests': {'bookId':1,'editionId':1,'note':'API test ebook borrow'},
    'favorite-books': {'bookId':1},
    'notifications': {'userId':1,'title':'API Test Notification','message':'Created by Postman API test','type':'SYSTEM'},
    'payments': {'borrowRecordId':1,'amount':50000,'method':'CASH','note':'API test payment'},
    'policies': {'name':'Policy API Test {{$timestamp}}','description':'Policy created by API test','value':'10','isActive':True},
    'reviews': {'bookId':1,'rating':5,'comment':'API test review'},
    'users': {'fullName':'User API Test {{$timestamp}}','email':'user_api_test_{{$timestamp}}@example.com','password':'Password@123','phoneNumber':'0900000000','address':'HCM','role':'READER'},
    'ai-summarize': {'text':'This is a sample long text used for AI summarize API test.'},
    'mail/test': {'email':'{{admin_email}}'},
}
query_samples = {
    'GET': 'page=1&limit=10&search=test&sortBy=createdAt&sortOrder=desc',
}

def endpoint_from_file(path: Path):
    rel = path.relative_to(API_ROOT).parent.as_posix()
    return '/api/' + rel.replace('[...path]', ':path*').replace('[bookId]', ':bookId').replace('[id]', ':id')

def route_key(path: Path):
    return path.relative_to(API_ROOT).parent.as_posix()

def methods_from_text(text):
    found = []
    for m in METHODS:
        if re.search(r'export\s+(?:async\s+function|const)\s+'+m+r'\b', text):
            found.append(m)
    return found

def desc_for(text, method, endpoint):
    pat = re.compile(r'//\s*'+method+r'\s+'+re.escape(endpoint.replace(':id','[id]').replace(':bookId','[bookId]').replace(':path*','[...path]'))+r'\s*-\s*(.+)')
    match = pat.search(text)
    if match:
        return match.group(1).strip()
    return f'{method} {endpoint}'

def auth_required(text, method):
    if 'requireAdmin' in text: return 'Admin token required'
    if 'requireLibrarian' in text: return 'Librarian/Admin token required'
    if 'requireAuth' in text: return 'Bearer token required'
    if 'optionalAuth' in text: return 'Optional Bearer token'
    return 'No token or public endpoint'

def choose_body(key, method):
    if method in ['GET','DELETE']: return ''
    if key in body_samples: return json.dumps(body_samples[key], ensure_ascii=False, indent=2)
    base = re.sub(r'/\[[^\]]+\]', '', key)
    return json.dumps(body_samples.get(base, {'sampleField':'sample value'}), ensure_ascii=False, indent=2)

def expected(method, endpoint):
    if method == 'POST': code='200/201'; data='created object / action result'
    elif method in ['PUT','PATCH']: code='200'; data='updated object / action result'
    elif method == 'DELETE': code='200'; data='null or deleted confirmation'
    else: code='200'; data='object/list payload'
    return f'HTTP {code}; response JSON có success=true, message hợp lệ, data={data}; lỗi trả về success=false/error message theo chuẩn handleRouteError.'

def postman_script(method, endpoint, positive=True):
    if positive:
        code_check = 'pm.expect(pm.response.code).to.be.oneOf([200,201]);'
        success = 'pm.expect(json.success).to.eql(true);'
    else:
        code_check = 'pm.expect(pm.response.code).to.be.oneOf([400,401,403,404,409,422,500]);'
        success = 'pm.expect(json.success).to.not.eql(true);'
    extra = ''
    if method == 'POST':
        extra = "\nif (json.data && json.data.id) pm.environment.set('last_created_id', json.data.id);"
    return "pm.test('Status code hợp lệ', function () {\n  " + code_check + "\n});\npm.test('Response là JSON và đúng contract', function () {\n  const json = pm.response.json();\n  pm.expect(json).to.be.an('object');\n  pm.expect(json).to.have.property('success');\n  " + success + "\n});" + extra

def db_note(method, endpoint):
    if method in ['POST','PUT','PATCH','DELETE']:
        return 'CheckDB: xác minh bản ghi thay đổi đúng trong MySQL/Prisma. Rollback: xóa bản ghi test theo id/code/email tạo bởi Postman hoặc khôi phục isDeleted=false nếu soft delete.'
    return 'CheckDB: đối chiếu dữ liệu response với bảng tương ứng qua Prisma/MySQL; không cần rollback vì chỉ đọc.'

routes=[]
for p in sorted(API_ROOT.rglob('route.ts')):
    text = p.read_text(encoding='utf-8', errors='ignore')
    ep = endpoint_from_file(p)
    key = route_key(p)
    for m in methods_from_text(text):
        routes.append({'method':m,'endpoint':ep,'key':key,'desc':desc_for(text,m,ep),'auth':auth_required(text,m),'file':str(p), 'text':text})

wb = Workbook()
ws = wb.active
ws.title = '1. API Inventory'
headers = ['No','Method','Endpoint','Mục đích API','Auth','Route file']
ws.append(headers)
for i,r in enumerate(routes,1): ws.append([i,r['method'],r['endpoint'],r['desc'],r['auth'],r['file']])

ws2 = wb.create_sheet('2. Test Cases')
headers2 = ['Test Case ID','API','Method','Mục đích test','Pre-condition','Input đầu vào Postman','Expected output','Postman test script','Check interface/DB','Rollback','Actual result','Pass/False','Ghi chú nếu False']
ws2.append(headers2)
case_no=1
for r in routes:
    ep=r['endpoint']; m=r['method']; key=r['key']
    url = '{{base_url}}' + ep.replace(':id','{{id}}').replace(':bookId','{{bookId}}').replace(':path*','{{file_path}}')
    if m=='GET' and '?' not in url and not any(x in ep for x in ['file','view']): url += '?' + query_samples['GET']
    input_obj = {'url':url,'method':m,'headers':{'Content-Type':'application/json','Authorization':'Bearer {{access_token}}' if 'token' in r['auth'].lower() or 'librarian' in r['auth'].lower() or 'admin' in r['auth'].lower() else ''},'body':choose_body(key,m)}
    ws2.append([f'TC-API-{case_no:03d}', f'{m} {ep}', m, f'Positive: {r["desc"]}', r['auth'], json.dumps(input_obj, ensure_ascii=False, indent=2), expected(m,ep), postman_script(m,ep,True), db_note(m,ep), 'Không rollback với GET; với POST/PUT/DELETE dùng id/email/code test để xóa hoặc khôi phục trạng thái.', 'Not executed in live Postman', 'PENDING', 'Chưa có server/database live để xác nhận runtime'])
    case_no+=1
    if m in ['POST','PUT','PATCH'] or ':id' in ep or ':bookId' in ep:
        neg_input = dict(input_obj); neg_input['body']='{}' if m in ['POST','PUT','PATCH'] else ''; neg_input['url']=url.replace('{{id}}','-1').replace('{{bookId}}','-1')
        ws2.append([f'TC-API-{case_no:03d}', f'{m} {ep}', m, f'Negative: validate input/permission/not found cho {r["desc"]}', r['auth'], json.dumps(neg_input, ensure_ascii=False, indent=2), 'HTTP 400/401/403/404/409; response JSON báo lỗi rõ ràng; DB không phát sinh dữ liệu rác.', postman_script(m,ep,False), 'CheckDB: không tạo/sửa/xóa sai dữ liệu khi request lỗi.', 'Không cần rollback nếu DB không đổi; nếu có dữ liệu test thì xóa theo khóa test.', 'Not executed in live Postman', 'PENDING', 'Cần chạy collection với environment thật để kết luận Pass/False'])
        case_no+=1

ws3 = wb.create_sheet('3. Postman Setup')
ws3.append(['Mục','Nội dung'])
ws3.append(['Environment variables','base_url=http://localhost:3000; access_token; refresh_token; admin_email; reader_email; reader_password; id; bookId; otp_code; file_path'])
ws3.append(['Pre-request login','Gọi POST {{base_url}}/api/auth/login, lưu accessToken: pm.environment.set("access_token", json.data.accessToken)'])
ws3.append(['Rollback guideline','Tạo dữ liệu với prefix/suffix API Test hoặc {{$timestamp}}, sau test xóa hard delete trong DB test hoặc gọi DELETE API rồi khôi phục nếu cần. Không chạy trên DB production.'])
ws3.append(['Interface check','Kiểm tra status code, response JSON contract, message, data, pagination, validation error hiển thị rõ trên Windows/Postman.'])
ws3.append(['Database check','Đối chiếu thay đổi bằng Prisma Studio/MySQL: user, book, author, category, borrow, payment, review, notification.'])

ws4 = wb.create_sheet('4. Execution Report')
ws4.append(['Hạng mục','Kết quả'])
ws4.append(['Ngày tạo báo cáo', datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')])
ws4.append(['Tổng API methods phát hiện', len(routes)])
ws4.append(['Tổng test cases sinh ra', case_no-1])
ws4.append(['Trạng thái thực thi Postman', 'PENDING - đã tạo testcase và script; cần chạy collection trên server/database test thật để cập nhật Pass/False'])
ws4.append(['Kết luận', 'Báo cáo bao phủ toàn bộ route.ts có method export trong src/app/api, gồm positive/negative, input/output, Postman test script, CheckDB và rollback guideline.'])

for sheet in wb.worksheets:
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='1F4E78')
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    for row in sheet.iter_rows():
        for cell in row:
            cell.border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
            cell.alignment = Alignment(vertical='top', wrap_text=True)
    for col in range(1, sheet.max_column+1):
        width = 18
        if sheet.title == '2. Test Cases': width = [16,28,12,45,28,55,45,60,45,45,22,12,45][col-1]
        elif col in [4,5,6]: width = 40
        sheet.column_dimensions[get_column_letter(col)].width = width
    sheet.freeze_panes = 'A2'

wb.save(OUT)
print(f'Generated {OUT}')
print(f'Routes: {len(routes)}, Test cases: {case_no-1}')
