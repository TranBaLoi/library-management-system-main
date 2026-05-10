# -*- coding: utf-8 -*-
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

path = r"C:\Users\Loi\Desktop\SQA\UnitTestReport_LibraryMS_Postman.xlsx"
wb = load_workbook(path)

thin_border = Border(left=Side(style="thin"), right=Side(style="thin"), top=Side(style="thin"), bottom=Side(style="thin"))
header_fill = PatternFill(start_color="D6E4F0", end_color="D6E4F0", fill_type="solid")
postman_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
cell_align = Alignment(vertical="top", wrap_text=True)

docker_cmd = "docker exec -it library-mysql mysql -u root -pyour_mysql_root_password library_management -e"

postman_auth = {
    "TC_AUTH_REG_01": f"Postman:\nPOST http://localhost:3000/api/auth/register\nBody (raw JSON):\n{{\"fullName\":\"Nguyen Van A\",\"email\":\"nguyenvana_postman@example.com\",\"password\":\"Password123!\",\"confirmPassword\":\"Password123!\"}}\n\nCheckDB (CMD):\n{docker_cmd} \"SELECT id,email,role,status FROM User WHERE email=\'nguyenvana_postman@example.com\';\"\nKỳ vọng: 1 dòng, role=READER, status=ACTIVE\n\nRollback:\n{docker_cmd} \"DELETE FROM User WHERE email=\'nguyenvana_postman@example.com\';\"",
    "TC_AUTH_REG_02": f"Postman:\nPOST http://localhost:3000/api/auth/register\nBody: {{\"fullName\":\"Test\",\"email\":\"invalid-email\",\"password\":\"Password123!\",\"confirmPassword\":\"Password123!\"}}\nKỳ vọng: HTTP 400\n\nCheckDB:\n{docker_cmd} \"SELECT COUNT(*) AS cnt FROM User WHERE email=\'invalid-email\';\"\nKỳ vọng: cnt=0 (không chèn gì)",
    "TC_AUTH_REG_03": f"Postman:\nPOST .../api/auth/register, password != confirmPassword\nKỳ vọng HTTP 400\n\nCheckDB:\n{docker_cmd} \"SELECT COUNT(*) AS cnt FROM User WHERE email=\'mismatch@test.com\';\"\nKỳ vọng: cnt=0",
    "TC_AUTH_REG_04": f"Chạy TC_AUTH_REG_01 trước, rồi POST lại cùng email.\nKỳ vọng HTTP 409\n\nCheckDB:\n{docker_cmd} \"SELECT COUNT(*) AS cnt FROM User WHERE email=\'nguyenvana_postman@example.com\';\"\nKỳ vọng: cnt=1 (không tăng)",
    "TC_AUTH_REG_05": f"Postman:\nPOST register, password=\'12345678\'\nKỳ vọng HTTP 400\n\nCheckDB:\n{docker_cmd} \"SELECT COUNT(*) AS cnt FROM User WHERE email=\'weak@test.com\';\"\nKỳ vọng: cnt=0",
    "TC_AUTH_REG_06": f"Postman:\nPOST register, fullName=\'\'\nKỳ vọng HTTP 400\n\nCheckDB:\n{docker_cmd} \"SELECT COUNT(*) AS cnt FROM User WHERE email=\'empty@test.com\';\"\nKỳ vọng: cnt=0",
    "TC_AUTH_REG_07": f"Postman:\nPOST register có phoneNumber=\'0901234567\'\n\nCheckDB:\n{docker_cmd} \"SELECT phoneNumber FROM User WHERE email=\'phone_ok@test.com\';\"\nKỳ vọng: 0901234567",
    "TC_AUTH_REG_08": f"Postman:\nPOST register, phoneNumber=\'123\'\nKỳ vọng HTTP 400\n\nCheckDB:\n{docker_cmd} \"SELECT COUNT(*) AS cnt FROM User WHERE email=\'phone_bad@test.com\';\"\nKỳ vọng: cnt=0",
    "TC_AUTH_LOGIN_01": f"Postman:\nPOST http://localhost:3000/api/auth/login\nBody: {{\"email\":\"nguyenvana_postman@example.com\",\"password\":\"Password123!\"}}\nKỳ vọng HTTP 200 có accessToken, refreshToken\n\nCheckDB:\n{docker_cmd} \"SELECT COUNT(*) AS cnt FROM RefreshToken WHERE userId=(SELECT id FROM User WHERE email=\'nguyenvana_postman@example.com\');\"\nKỳ vọng: cnt >= 1",
    "TC_AUTH_LOGIN_02": f"Postman:\nPOST login, email không tồn tại\nKỳ vọng HTTP 401\n\nCheckDB: không có RefreshToken mới cho email này",
    "TC_AUTH_LOGIN_03": f"Postman:\nPOST login, password sai\nKỳ vọng HTTP 401\n\nCheckDB: số RefreshToken không tăng",
    "TC_AUTH_LOGIN_04": f"Setup DB trước:\n{docker_cmd} \"UPDATE User SET status=\'INACTIVE\' WHERE email=\'nguyenvana_postman@example.com\';\"\n\nPostman: POST login\nKỳ vọng HTTP 401\n\nRollback:\n{docker_cmd} \"UPDATE User SET status=\'ACTIVE\' WHERE email=\'nguyenvana_postman@example.com\';\"",
    "TC_AUTH_LOGIN_05": f"Setup DB:\n{docker_cmd} \"UPDATE User SET isDeleted=1 WHERE email=\'nguyenvana_postman@example.com\';\"\n\nPostman: POST login\nKỳ vọng HTTP 401\n\nRollback:\n{docker_cmd} \"UPDATE User SET isDeleted=0 WHERE email=\'nguyenvana_postman@example.com\';\"",
    "TC_AUTH_LOGIN_06": "Postman:\nPOST login, email=\'not_an_email\'\nKỳ vọng HTTP 400\n\nKhông cần checkDB",
    "TC_AUTH_LOGIN_07": "Postman:\nPOST login, password=\'\'\nKỳ vọng HTTP 400\n\nKhông cần checkDB",
    "TC_AUTH_LOGIN_08": f"Setup DB:\n{docker_cmd} \"UPDATE User SET firstLoginAt=NULL WHERE email=\'nguyenvana_postman@example.com\';\"\n\nPostman: POST login\n\nCheckDB:\n{docker_cmd} \"SELECT firstLoginAt FROM User WHERE email=\'nguyenvana_postman@example.com\';\"\nKỳ vọng: firstLoginAt != NULL",
    "TC_AUTH_LOGIN_09": f"Postman: POST login (firstLoginAt đã có giá trị)\n\nCheckDB:\n{docker_cmd} \"SELECT firstLoginAt FROM User WHERE email=\'nguyenvana_postman@example.com\';\"\nKỳ vọng: firstLoginAt không đổi",
    "TC_AUTH_LOGIN_10": f"Postman:\nPOST login với rememberMe=true\n\nCheckDB:\n{docker_cmd} \"SELECT expiresAt FROM RefreshToken WHERE userId=(SELECT id FROM User WHERE email=\'nguyenvana_postman@example.com\') ORDER BY createdAt DESC LIMIT 1;\"\nKỳ vọng: expiresAt ~ +30 ngày",
    "TC_AUTH_REFRESH_01": "Postman:\nPOST http://localhost:3000/api/auth/refresh\nCookie: refreshToken=<token từ login>\nKỳ vọng HTTP 200 có accessToken mới",
    "TC_AUTH_REFRESH_02": "Postman:\nPOST /api/auth/refresh, cookie token giả\nKỳ vọng HTTP 401",
    "TC_AUTH_REFRESH_03": f"Setup DB:\n{docker_cmd} \"UPDATE RefreshToken SET expiresAt=NOW()-INTERVAL 1 DAY WHERE userId=(SELECT id FROM User WHERE email=\'nguyenvana_postman@example.com\') LIMIT 1;\"\n\nPostman: POST /api/auth/refresh\nKỳ vọng HTTP 401\n\nCheckDB:\n{docker_cmd} \"SELECT COUNT(*) FROM RefreshToken WHERE userId=(SELECT id FROM User WHERE email=\'nguyenvana_postman@example.com\');\"\nKỳ vọng: token đó đã bị xóa",
    "TC_AUTH_REFRESH_04": f"Setup: set user INACTIVE\nPostman: POST /api/auth/refresh\nKỳ vọng HTTP 401\nRollback: set ACTIVE lại",
    "TC_AUTH_CHANGEPWD_01": f"Postman:\nPOST http://localhost:3000/api/auth/change-password\nHeaders: Cookie accessToken + refreshToken\nBody: {{\"currentPassword\":\"Password123!\",\"newPassword\":\"NewPass@456\",\"confirmNewPassword\":\"NewPass@456\"}}\n\nCheckDB:\n{docker_cmd} \"SELECT password FROM User WHERE email=\'nguyenvana_postman@example.com\';\"\nKỳ vọng: hash đổi\n\n{docker_cmd} \"SELECT COUNT(*) FROM RefreshToken WHERE userId=(SELECT id FROM User WHERE email=\'nguyenvana_postman@example.com\');\"\nKỳ vọng: 0 (logout all)",
    "TC_AUTH_CHANGEPWD_02": "Postman:\nPOST change-password, currentPassword sai\nKỳ vọng HTTP 400\nCheckDB: hash không đổi",
    "TC_AUTH_CHANGEPWD_03": "Postman:\nPOST change-password, newPassword != confirm\nKỳ vọng HTTP 400\nCheckDB: hash không đổi",
    "TC_AUTH_CHANGEPWD_04": "Postman:\nPOST change-password, newPassword == currentPassword\nKỳ vọng HTTP 400\nCheckDB: hash không đổi",
    "TC_AUTH_CHANGEPWD_05": "Không test trực tiếp bằng Postman (cần token user không tồn tại)",
    "TC_AUTH_LOGOUT_01": f"Postman:\nPOST http://localhost:3000/api/auth/logout\nCookie: refreshToken=<token>\nKỳ vọng HTTP 200\n\nCheckDB:\n{docker_cmd} \"SELECT COUNT(*) FROM RefreshToken WHERE token=\'<refreshToken cũ>\';\"\nKỳ vọng: 0 (token đã xóa)",
    "TC_AUTH_LOGOUT_02": f"logoutAll chỉ gọi nội bộ sau changePassword\nCheckDB qua TC_AUTH_CHANGEPWD_01:\n{docker_cmd} \"SELECT COUNT(*) FROM RefreshToken WHERE userId=<id>;\"\nKỳ vọng: 0",
    "TC_AUTH_CLEANUP_01": f"Service nội bộ/cron, không có endpoint\nCheckDB trước:\n{docker_cmd} \"SELECT COUNT(*) FROM RefreshToken WHERE expiresAt < NOW();\"\nSau khi chạy cron/script: count giảm",
    "TC_AUTH_CLEANUP_02": f"CheckDB:\n{docker_cmd} \"SELECT COUNT(*) FROM RefreshToken WHERE expiresAt < NOW();\"\nKỳ vọng: 0 nếu không có token hết hạn",
}

postman_book = {
    "TC_BOOK_TRANSFORM_01": f"Postman:\nGET http://localhost:3000/api/books?page=1&limit=5\nKỳ vọng response có categories, bookItemsCount, averageRating\n\nCheckDB:\n{docker_cmd} \"SELECT b.id,b.title,(SELECT COUNT(*) FROM BookItem WHERE bookId=b.id AND isDeleted=0) AS itemCount,(SELECT AVG(rating) FROM Review WHERE bookId=b.id AND isDeleted=0) AS avgRating FROM Book b WHERE b.isDeleted=0 LIMIT 5;\"",
    "TC_BOOK_TRANSFORM_02": f"Postman: GET /api/books\nChọn book chưa có review\n\nCheckDB:\n{docker_cmd} \"SELECT b.id FROM Book b WHERE b.id NOT IN (SELECT DISTINCT bookId FROM Review WHERE isDeleted=0) LIMIT 1;\"\nKỳ vọng averageRating=0",
    "TC_BOOK_TRANSFORM_03": f"Postman: GET /api/books\nChọn book không có category/edition/item\n\nCheckDB:\n{docker_cmd} \"SELECT b.id FROM Book b WHERE b.id NOT IN (SELECT bookId FROM BookCategory) AND b.isDeleted=0 LIMIT 1;\"\nKỳ vọng categories=[]",
    "TC_BOOK_FILTER_01": f"Postman: GET /api/books\nKỳ vọng: response không chứa sách isDeleted=true\n\nCheckDB:\n{docker_cmd} \"SELECT COUNT(*) FROM Book WHERE isDeleted=1;\"",
    "TC_BOOK_FILTER_02": f"Postman: GET /api/books?authorIds=1&authorIds=2\n\nCheckDB:\n{docker_cmd} \"SELECT id,authorId FROM Book WHERE authorId IN (1,2) AND isDeleted=0;\"",
    "TC_BOOK_FILTER_03": f"Postman: GET /api/books?categoryIds=5&categoryIds=6\n\nCheckDB:\n{docker_cmd} \"SELECT bc.bookId FROM BookCategory bc WHERE bc.categoryId IN (5,6);\"",
    "TC_BOOK_FILTER_04": f"Postman: GET /api/books?languageCodes=vi&languageCodes=en\n\nCheckDB:\n{docker_cmd} \"SELECT id,language FROM Book WHERE language IN (\'vi\',\'en\') AND isDeleted=0;\"",
    "TC_BOOK_FILTER_05": f"Postman: GET /api/books?publishYearFrom=2000&publishYearTo=2020\n\nCheckDB:\n{docker_cmd} \"SELECT id,publishYear FROM Book WHERE publishYear BETWEEN 2000 AND 2020 AND isDeleted=0;\"",
    "TC_BOOK_FILTER_06": f"Postman: GET /api/books?availableAt=ebook\n\nCheckDB:\n{docker_cmd} \"SELECT DISTINCT bookId FROM BookEdition WHERE format=\'EBOOK\' AND isDeleted=0;\"",
    "TC_BOOK_FILTER_07": f"Postman: GET /api/books?availableAt=book-copy\n\nCheckDB:\n{docker_cmd} \"SELECT DISTINCT bookId FROM BookItem WHERE status=\'AVAILABLE\' AND isDeleted=0;\"",
    "TC_BOOK_FILTER_08": f"Postman: GET /api/books?availableAt=ebook&availableAt=book-copy\n\nCheckDB: bookId phải tồn tại cả trong BookEdition(EBOOK) và BookItem(AVAILABLE)",
    "TC_BOOK_SORT_01": f"Postman: GET /api/books\nKỳ vọng createdAt desc\n\nCheckDB:\n{docker_cmd} \"SELECT id,createdAt FROM Book WHERE isDeleted=0 ORDER BY createdAt DESC LIMIT 10;\"",
    "TC_BOOK_SORT_02": f"Postman: GET /api/books?sortBy=title&sortOrder=asc\n\nCheckDB:\n{docker_cmd} \"SELECT title FROM Book WHERE isDeleted=0 ORDER BY title ASC LIMIT 10;\"",
    "TC_BOOK_SORT_03": f"Postman: GET /api/books?sortBy=price&sortOrder=desc\n\nCheckDB:\n{docker_cmd} \"SELECT price FROM Book WHERE isDeleted=0 ORDER BY price DESC LIMIT 10;\"",
    "TC_BOOK_SORT_04": "Postman: GET /api/books?sortBy=unknownField&sortOrder=asc\nKỳ vọng fallback createdAt desc\nSo sánh thứ tự với TC_BOOK_SORT_01",
    "TC_BOOK_LIST_01": f"Postman: GET /api/books?page=1&limit=10\nKỳ vọng books.length<=10, pagination.total khớp\n\nCheckDB:\n{docker_cmd} \"SELECT COUNT(*) FROM Book WHERE isDeleted=0;\"",
    "TC_BOOK_LIST_02": f"Postman: GET /api/books?page=3&limit=20\nKỳ vọng offset=40\n\nCheckDB:\n{docker_cmd} \"SELECT id,title FROM Book WHERE isDeleted=0 ORDER BY createdAt DESC LIMIT 20 OFFSET 40;\"",
    "TC_BOOK_LIST_03": f"Postman: GET /api/books?search=clean\n\nCheckDB:\n{docker_cmd} \"SELECT b.id,b.title FROM Book b LEFT JOIN Author a ON b.authorId=a.id WHERE b.title LIKE \'%clean%\' OR b.isbn LIKE \'%clean%\' OR b.publisher LIKE \'%clean%\' OR b.description LIKE \'%clean%\' OR a.fullName LIKE \'%clean%\';\"",
    "TC_BOOK_LIST_04": "Postman: GET /api/books?search=clean&availableAt=ebook&availableAt=book-copy\nKỳ vọng vừa match keyword vừa thỏa cả 2 availability\nCheckDB kết hợp LIKE + ebook + book-copy",
    "TC_BOOK_LIST_05": f"Postman: GET /api/books?sortBy=publishYear&sortOrder=desc\n\nCheckDB:\n{docker_cmd} \"SELECT publishYear FROM Book WHERE isDeleted=0 ORDER BY publishYear DESC LIMIT 10;\"",
    "TC_BOOK_LIST_06": f"Postman: GET /api/books?authorIds=1&categoryIds=2&languageCodes=vi\n\nCheckDB:\n{docker_cmd} \"SELECT b.id,b.title FROM Book b JOIN BookCategory bc ON b.id=bc.bookId WHERE b.authorId=1 AND bc.categoryId=2 AND b.language=\'vi\' AND b.isDeleted=0;\"",
}

for sheet_name, mapping in [("3.TC_XacThuc", postman_auth), ("4.TC_TimKiemSach", postman_book)]:
    if sheet_name not in wb.sheetnames:
        continue
    ws = wb[sheet_name]
    # Check if Postman column already exists
    postman_col = None
    for c in range(1, ws.max_column + 1):
        val = ws.cell(row=2, column=c).value
        if val and "Postman" in str(val):
            postman_col = c
            break
    if not postman_col:
        postman_col = ws.max_column + 1
        ws.cell(row=2, column=postman_col, value="L\u1ec7nh Postman & CheckDB (Docker)")
        ws.cell(row=2, column=postman_col).font = Font(bold=True)
        ws.cell(row=2, column=postman_col).fill = header_fill
        ws.cell(row=2, column=postman_col).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.cell(row=2, column=postman_col).border = thin_border
    ws.column_dimensions[get_column_letter(postman_col)].width = 80

    for row in range(3, ws.max_row + 1):
        tc_id = ws.cell(row=row, column=1).value
        if tc_id and str(tc_id) in mapping:
            c = ws.cell(row=row, column=postman_col, value=mapping[str(tc_id)])
            c.alignment = cell_align
            c.border = thin_border
            c.fill = postman_fill

wb.save(path)
print(f"DONE: {path}")