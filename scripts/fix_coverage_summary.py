from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

path = r"C:\Users\Loi\Desktop\SQA\UnitTestReport_LibraryMS_Postman.xlsx"
wb = load_workbook(path)
high_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
thin_border = Border(left=Side(style="thin"), right=Side(style="thin"), top=Side(style="thin"), bottom=Side(style="thin"))
cell_align = Alignment(vertical="top", wrap_text=True)
center_align = Alignment(horizontal="center", vertical="top", wrap_text=True)

ws6 = wb["6.Code Coverage"]

# Sửa dòng tóm tắt auth.service.ts
for row in range(1, ws6.max_row + 1):
    val = ws6.cell(row=row, column=1).value
    if val == "auth.service.ts":
        ws6.cell(row=row, column=2, value="100% Stmts").fill = high_fill
        ws6.cell(row=row, column=3, value="96.72% Branch").fill = high_fill
        ws6.cell(row=row, column=4, value="100% Funcs").fill = high_fill
        ws6.cell(row=row, column=5, value="100% Lines").fill = high_fill
        for c in range(1, 6):
            ws6.cell(row=row, column=c).border = thin_border
            ws6.cell(row=row, column=c).alignment = center_align

    if val and "Tổng kết" in str(val):
        ws6.cell(row=row, column=2, value="97.45%").fill = high_fill
        ws6.cell(row=row, column=3, value="93.10%").fill = high_fill
        ws6.cell(row=row, column=4, value="100%").fill = high_fill
        ws6.cell(row=row, column=5, value="97.41%").fill = high_fill
        ws6.cell(row=row, column=6, value="Coverage All files sau khi chỉ đo 2 service chính + 54 tests")
        for c in range(1, 7):
            ws6.cell(row=row, column=c).border = thin_border
            ws6.cell(row=row, column=c).alignment = center_align if c <= 5 else cell_align

# Cập nhật dòng auth.service.ts (dòng r1) trong tóm tắt

wb.save(path)
print("DONE: updated coverage summary")