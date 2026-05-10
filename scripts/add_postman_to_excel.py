# -*- coding: utf-8 -*-
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

path = r"C:\Users\Loi\Desktop\SQA\UnitTestReport_LibraryMS.xlsx"
wb = load_workbook(path)

thin_border = Border(left=Side(style="thin"), right=Side(style="thin"), top=Side(style="thin"), bottom=Side(style="thin"))
header_fill = PatternFill(start_color="D6E4F0", end_color="D6E4F0", fill_type="solid")
postman_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
section_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
title_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
cell_align = Alignment(vertical="top", wrap_text=True)
center_align = Alignment(horizontal="center", vertical="top", wrap_text=True)

base_url = "http://localhost:3000"

postman_auth = {
    "TC_AUTH_REG_01": "POST {{baseUrl}}/api/auth/register\nBody JSON:\n{\n  \"fullName\": \"Nguyen Van A\",\n  \"email\": \"nguyenvana_postman@example.com\",\n  \"password\": \"Password123!\",\n  \"confirmPassword\": \"Password123!\",\n  \"phoneNumber\": \"0901234567\"\n}\n\nCheckDB SQL:\nSELECT id,email,fullName,role,status,isDeleted FROM User WHERE email='nguyenvana_postman@example.com';\nKỳ vọng: có 1 dòng user ACTIVE, role READER.",
    "TC_AUTH_REG_02": "POST {{baseUrl}}/api/auth/register với email='invalid-email'\nKỳ vọng HTTP 400.\n\nCheckDB SQL:\nSELECT COUNT(*) FROM User WHERE email='invalid-email';\nKỳ vọng: 0, không chèn dữ liệu.",
    "TC_AUTH_REG_03": "POST {{baseUrl}}/api/auth/register với password != confirmPassword\nKỳ vọng HTTP 400.\n\nCheckDB SQL:\nSELECT COUNT(*) FROM User WHERE email='password_mismatch@example.com';\nKỳ vọng: 0.",
    "TC_AUTH_REG_04": "Chạy TC_AUTH_REG_01 trước để tạo user, sau đó POST lại cùng email.\nKỳ vọng HTTP 409.\n\nCheckDB SQL:\nSELECT COUNT(*) FROM User WHERE email='nguyenvana_postman@example.com';\nKỳ vọng: vẫn = 1, không tạo trùng.",
    "TC_AUTH_REG_05": "POST {{baseUrl}}/api/auth/register với password='12345678'\nKỳ vọng HTTP 400.\n\nCheckDB SQL:\nSELECT COUNT(*) FROM User WHERE email='weak_password@example.com';\nKỳ vọng: 0.",
    "TC_AUTH_REG_06": "POST {{baseUrl}}/api/auth/register với fullName=''\nKỳ vọng HTTP 400.\n\nCheckDB SQL:\nSELECT COUNT(*) FROM User WHERE email='empty_name@example.com';\nKỳ vọng: 0.",
    "TC_AUTH_REG_07": "POST {{baseUrl}}/api/auth/register có phoneNumber='0901234567'\n\nCheckDB SQL:\nSELECT phoneNumber FROM User WHERE email='phone_valid@example.com';\nKỳ vọng: phoneNumber='0901234567'.",
    "TC_AUTH_REG_08": "POST {{baseUrl}}/api/auth/register với phoneNumber='123'\nKỳ vọng HTTP 400.\n\nCheckDB SQL:\nSELECT COUNT(*) FROM User WHERE email='phone_invalid@example.com';\nKỳ vọng: 0.",
    "TC_AUTH_LOGIN_01": "POST {{baseUrl}}/api/auth/login\nBody JSON:\n{\"email\":\"nguyenvana_postman@example.com\",\"password\":\"Password123!\"}\n\nCheckDB SQL:\nSELECT * FROM RefreshToken WHERE userId=(SELECT id FROM User WHERE email='nguyenvana_postman@example.com');\nKỳ vọng: có refresh token mới.",
    "TC_AUTH_LOGIN_02": "POST {{baseUrl}}/api/auth/login với email không tồn tại.\nKỳ vọng HTTP 401.\n\nCheckDB SQL:\nSELECT COUNT(*) FROM RefreshToken WHERE userId IS NULL;\nKỳ vọng: không phát sinh token hợp lệ.",
    "TC_AUTH_LOGIN_03": "POST {{baseUrl}}/api/auth/login với password sai.\nKỳ vọng HTTP 401.\n\nCheckDB SQL:\nSELECT COUNT(*) FROM RefreshToken WHERE userId=(SELECT id FROM User WHERE email='nguyenvana_postman@example.com');\nKỳ vọng: không tăng số token so với trước request.",
    "TC_AUTH_LOGIN_04": "Set DB user status=INACTIVE rồi POST /api/auth/login.\n\nSQL setup:\nUPDATE User SET status='INACTIVE' WHERE email='inactive_user@example.com';\nKỳ vọng HTTP 401.\nRollback:\nUPDATE User SET status='ACTIVE' WHERE email='inactive_user@example.com';",
    "TC_AUTH_LOGIN_05": "Set DB user isDeleted=true rồi POST /api/auth/login.\n\nSQL setup:\nUPDATE User SET isDeleted=1 WHERE email='deleted_user@example.com';\nKỳ vọng HTTP 401.\nRollback:\nUPDATE User SET isDeleted=0 WHERE email='deleted_user@example.com';",
    "TC_AUTH_LOGIN_06": "POST {{baseUrl}}/api/auth/login với email='not_an_email'.\nKỳ vọng HTTP 400.\n\nCheckDB: không có RefreshToken mới.",
    "TC_AUTH_LOGIN_07": "POST {{baseUrl}}/api/auth/login với password=''.\nKỳ vọng HTTP 400.\n\nCheckDB: không có RefreshToken mới.",
    "TC_AUTH_LOGIN_08": "Trước login:\nUPDATE User SET firstLoginAt=NULL WHERE email='nguyenvana_postman@example.com';\nPOST /api/auth/login.\nCheckDB:\nSELECT firstLoginAt FROM User WHERE email='nguyenvana_postman@example.com';\nKỳ vọng: firstLoginAt không NULL.",
    "TC_AUTH_LOGIN_09": "Trước login: user đã có firstLoginAt.\nPOST /api/auth/login.\nCheckDB:\nSELECT firstLoginAt FROM User WHERE email='nguyenvana_postman@example.com';\nKỳ vọng: giá trị không bị reset bất thường.",
    "TC_AUTH_LOGIN_10": "POST /api/auth/login với rememberMe=true.\nBody JSON: {\"email\":\"...\",\"password\":\"Password123!\",\"rememberMe\":true}\nCheckDB:\nSELECT expiresAt FROM RefreshToken ORDER BY expiresAt DESC LIMIT 1;\nKỳ vọng: expiresAt khoảng +30 ngày.",
    "TC_AUTH_REFRESH_01": "POST {{baseUrl}}/api/auth/refresh với cookie refreshToken từ login.\nKỳ vọng HTTP 200, có accessToken mới.\nCheckDB: RefreshToken vẫn tồn tại.",
    "TC_AUTH_REFRESH_02": "Gửi cookie refreshToken giả/đã xóa.\nKỳ vọng HTTP 401.\nCheckDB: không tạo token mới.",
    "TC_AUTH_REFRESH_03": "SQL setup token hết hạn:\nUPDATE RefreshToken SET expiresAt=NOW() - INTERVAL 1 DAY WHERE id='<tokenId>';\nPOST /api/auth/refresh.\nCheckDB:\nSELECT COUNT(*) FROM RefreshToken WHERE id='<tokenId>';\nKỳ vọng: 0, token đã bị xóa.",
    "TC_AUTH_REFRESH_04": "Set user INACTIVE rồi POST /api/auth/refresh với token user đó.\nKỳ vọng HTTP 401.\nRollback: set ACTIVE lại.",
    "TC_AUTH_CHANGEPWD_01": "POST {{baseUrl}}/api/auth/change-password với Bearer token/cookie auth.\nBody JSON:\n{\"currentPassword\":\"Password123!\",\"newPassword\":\"NewPass@456\",\"confirmNewPassword\":\"NewPass@456\"}\nCheckDB:\nSELECT password FROM User WHERE email='nguyenvana_postman@example.com';\nKỳ vọng: hash password thay đổi.\nSELECT COUNT(*) FROM RefreshToken WHERE userId=<id>; Kỳ vọng: 0.",
    "TC_AUTH_CHANGEPWD_02": "POST /api/auth/change-password với currentPassword sai.\nKỳ vọng HTTP 400.\nCheckDB: password hash không đổi.",
    "TC_AUTH_CHANGEPWD_03": "POST /api/auth/change-password với newPassword != confirm.\nKỳ vọng HTTP 400.\nCheckDB: password hash không đổi.",
    "TC_AUTH_CHANGEPWD_04": "POST /api/auth/change-password với newPassword giống mật khẩu hiện tại.\nKỳ vọng HTTP 400.\nCheckDB: password hash không đổi.",
    "TC_AUTH_CHANGEPWD_05": "Không test trực tiếp bằng Postman nếu không có userId ảo qua API.\nCó thể test integration bằng token user đã bị xóa.\nCheckDB: không phát sinh update User.",
    "TC_AUTH_LOGOUT_01": "POST {{baseUrl}}/api/auth/logout với cookie refreshToken.\nCheckDB:\nSELECT COUNT(*) FROM RefreshToken WHERE token='<refreshToken>';\nKỳ vọng: 0, token đã xóa.",
    "TC_AUTH_LOGOUT_02": "Không có API riêng logoutAll; được gọi sau đổi mật khẩu.\nCheck qua TC_AUTH_CHANGEPWD_01:\nSELECT COUNT(*) FROM RefreshToken WHERE userId=<id>;\nKỳ vọng: 0.",
    "TC_AUTH_CLEANUP_01": "cleanupExpiredTokens là service/cron, không có endpoint public.\nCó thể kiểm qua script hoặc endpoint nội bộ nếu có.\nSQL trước/sau:\nSELECT COUNT(*) FROM RefreshToken WHERE expiresAt < NOW();",
    "TC_AUTH_CLEANUP_02": "Khi không có token hết hạn:\nSELECT COUNT(*) FROM RefreshToken WHERE expiresAt < NOW();\nKỳ vọng: 0. Service trả count=0.",
}

postman_book = {
    "TC_BOOK_TRANSFORM_01": "GET {{baseUrl}}/api/books?page=1&limit=10\nKỳ vọng response book có categories, bookItemsCount, bookEbookCount, bookAudioCount, averageRating.\nCheckDB SQL: kiểm tra Book, BookCategory, BookEdition, Review tương ứng.",
    "TC_BOOK_TRANSFORM_02": "Tạo/kiểm sách chưa có Review rồi GET /api/books.\nCheckDB:\nSELECT COUNT(*) FROM Review WHERE bookId=<bookId> AND isDeleted=0;\nKỳ vọng API averageRating=0.",
    "TC_BOOK_TRANSFORM_03": "GET /api/books với sách không có category/edition/item.\nCheckDB các bảng liên quan count=0.\nKỳ vọng API categories=[], count=0.",
    "TC_BOOK_FILTER_01": "GET {{baseUrl}}/api/books\nCheckDB SQL:\nSELECT COUNT(*) FROM Book WHERE isDeleted=1;\nKỳ vọng response không chứa sách isDeleted=true.",
    "TC_BOOK_FILTER_02": "GET {{baseUrl}}/api/books?authorIds=1&authorIds=2\nCheckDB:\nSELECT DISTINCT authorId FROM Book WHERE authorId IN (1,2);\nKỳ vọng response chỉ có authorId 1 hoặc 2.",
    "TC_BOOK_FILTER_03": "GET {{baseUrl}}/api/books?categoryIds=5&categoryIds=6\nCheckDB:\nSELECT * FROM BookCategory WHERE categoryId IN (5,6);\nKỳ vọng response chỉ gồm sách thuộc category 5/6.",
    "TC_BOOK_FILTER_04": "GET {{baseUrl}}/api/books?languageCodes=vi&languageCodes=en\nCheckDB:\nSELECT DISTINCT language FROM Book WHERE language IN ('vi','en');\nKỳ vọng response language chỉ vi/en.",
    "TC_BOOK_FILTER_05": "GET {{baseUrl}}/api/books?publishYearFrom=2000&publishYearTo=2020\nCheckDB:\nSELECT id,title,publishYear FROM Book WHERE publishYear BETWEEN 2000 AND 2020;\nKỳ vọng response publishYear trong khoảng.",
    "TC_BOOK_FILTER_06": "GET {{baseUrl}}/api/books?availableAt=ebook\nCheckDB:\nSELECT DISTINCT bookId FROM BookEdition WHERE format='EBOOK' AND isDeleted=0;\nKỳ vọng response chỉ gồm sách có EBOOK.",
    "TC_BOOK_FILTER_07": "GET {{baseUrl}}/api/books?availableAt=book-copy\nCheckDB:\nSELECT DISTINCT bookId FROM BookItem WHERE status='AVAILABLE' AND isDeleted=0;\nKỳ vọng response chỉ gồm sách có bản in sẵn sàng.",
    "TC_BOOK_FILTER_08": "GET {{baseUrl}}/api/books?availableAt=ebook&availableAt=book-copy\nCheckDB: bookId phải tồn tại ở cả BookEdition(EBOOK) và BookItem(AVAILABLE).\nKỳ vọng response chỉ gồm sách thỏa cả 2.",
    "TC_BOOK_SORT_01": "GET {{baseUrl}}/api/books\nKỳ vọng sắp xếp createdAt giảm dần.\nCheckDB: SELECT id,createdAt FROM Book WHERE isDeleted=0 ORDER BY createdAt DESC;",
    "TC_BOOK_SORT_02": "GET {{baseUrl}}/api/books?sortBy=title&sortOrder=asc\nCheckDB:\nSELECT title FROM Book WHERE isDeleted=0 ORDER BY title ASC;",
    "TC_BOOK_SORT_03": "GET {{baseUrl}}/api/books?sortBy=price&sortOrder=desc\nCheckDB:\nSELECT price FROM Book WHERE isDeleted=0 ORDER BY price DESC;",
    "TC_BOOK_SORT_04": "GET {{baseUrl}}/api/books?sortBy=unknownField&sortOrder=asc\nKỳ vọng fallback createdAt desc.\nCheckDB giống TC_BOOK_SORT_01.",
    "TC_BOOK_LIST_01": "GET {{baseUrl}}/api/books?page=1&limit=10\nKỳ vọng pagination.page=1, limit=10, books.length<=10.\nCheckDB: SELECT COUNT(*) FROM Book WHERE isDeleted=0;",
    "TC_BOOK_LIST_02": "GET {{baseUrl}}/api/books?page=3&limit=20\nKỳ vọng response là trang thứ 3.\nCheckDB tương đương OFFSET 40 LIMIT 20.",
    "TC_BOOK_LIST_03": "GET {{baseUrl}}/api/books?search=clean\nCheckDB:\nSELECT * FROM Book b LEFT JOIN Author a ON b.authorId=a.id WHERE b.title LIKE '%clean%' OR b.isbn LIKE '%clean%' OR b.publisher LIKE '%clean%' OR b.description LIKE '%clean%' OR a.fullName LIKE '%clean%';",
    "TC_BOOK_LIST_04": "GET {{baseUrl}}/api/books?search=clean&availableAt=ebook&availableAt=book-copy\nKỳ vọng response vừa match keyword vừa có ebook + bản in.\nCheckDB kết hợp truy vấn ở TC_BOOK_LIST_03 và TC_BOOK_FILTER_08.",
    "TC_BOOK_LIST_05": "GET {{baseUrl}}/api/books?sortBy=publishYear&sortOrder=desc\nCheckDB:\nSELECT publishYear FROM Book WHERE isDeleted=0 ORDER BY publishYear DESC;",
    "TC_BOOK_LIST_06": "GET {{baseUrl}}/api/books?authorIds=1&categoryIds=2&languageCodes=vi\nCheckDB: Book.authorId=1, BookCategory.categoryId=2, Book.language='vi'.\nKỳ vọng response thỏa cả 3.",
}

# Add Postman column to TC sheets
for sheet_name, mapping in [("3.TC_XacThuc", postman_auth), ("4.TC_TimKiemSach", postman_book)]:
    ws = wb[sheet_name]
    new_col = ws.max_column + 1
    ws.cell(row=2, column=new_col, value="Lệnh Postman / CheckDB SQL")
    ws.cell(row=2, column=new_col).font = Font(bold=True)
    ws.cell(row=2, column=new_col).fill = header_fill
    ws.cell(row=2, column=new_col).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.cell(row=2, column=new_col).border = thin_border
    ws.column_dimensions[get_column_letter(new_col)].width = 70
    for row in range(3, ws.max_row + 1):
        tc_id = ws.cell(row=row, column=1).value
        if tc_id in mapping:
            c = ws.cell(row=row, column=new_col, value=mapping[tc_id])
            c.alignment = cell_align
            c.border = thin_border
            c.fill = postman_fill

# Create or replace Postman guide sheet
if "8.Postman_CheckDB" in wb.sheetnames:
    del wb["8.Postman_CheckDB"]
ws = wb.create_sheet("8.Postman_CheckDB")
for col, width in {"A": 8, "B": 28, "C": 42, "D": 65, "E": 55, "F": 35}.items():
    ws.column_dimensions[col].width = width

ws.merge_cells("A1:F1")
ws["A1"] = "HƯỚNG DẪN POSTMAN & CHECKDB SAU KHI THỰC THI API"
ws["A1"].font = Font(bold=True, size=14, color="FFFFFF")
ws["A1"].fill = title_fill
ws["A1"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

headers = ["STT", "Nhóm", "Mục tiêu", "Postman Request", "CheckDB SQL / Kiểm tra DB", "Rollback / Cleanup"]
for i, h in enumerate(headers, 1):
    cell = ws.cell(row=2, column=i, value=h)
    cell.font = Font(bold=True)
    cell.fill = header_fill
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = thin_border

rows = [
    [1, "Auth - Register", "Kiểm tra dữ liệu User đã được chèn", "POST {{baseUrl}}/api/auth/register\nBody JSON user test", "SELECT id,email,role,status FROM User WHERE email='nguyenvana_postman@example.com';\nKỳ vọng: có 1 dòng ACTIVE/READER", "DELETE FROM RefreshToken WHERE userId=(SELECT id FROM User WHERE email='nguyenvana_postman@example.com');\nDELETE FROM User WHERE email='nguyenvana_postman@example.com';"],
    [2, "Auth - Login", "Kiểm tra RefreshToken đã được tạo", "POST {{baseUrl}}/api/auth/login\nBody JSON email/password đúng", "SELECT * FROM RefreshToken WHERE userId=(SELECT id FROM User WHERE email='nguyenvana_postman@example.com');\nKỳ vọng: có token mới", "DELETE FROM RefreshToken WHERE userId=<userId>;"],
    [3, "Auth - Refresh", "Kiểm tra refresh token còn hợp lệ", "POST {{baseUrl}}/api/auth/refresh\nCookie: refreshToken=<token>", "SELECT token,expiresAt FROM RefreshToken WHERE token='<token>';\nKỳ vọng: vẫn tồn tại nếu chưa expired", "Không cần rollback nếu chỉ đọc/làm mới accessToken"],
    [4, "Auth - Logout", "Kiểm tra token đã bị xóa", "POST {{baseUrl}}/api/auth/logout\nCookie: refreshToken=<token>", "SELECT COUNT(*) FROM RefreshToken WHERE token='<token>';\nKỳ vọng: 0", "Không cần rollback vì hành vi mong muốn là xóa token"],
    [5, "Auth - Change Password", "Kiểm tra password hash đổi và token bị xóa", "POST {{baseUrl}}/api/auth/change-password\nAuthorization/Cookie hợp lệ\nBody current/new/confirm", "SELECT password FROM User WHERE email='nguyenvana_postman@example.com';\nSELECT COUNT(*) FROM RefreshToken WHERE userId=<userId>;\nKỳ vọng: hash đổi, token count=0", "Đổi lại mật khẩu cũ bằng API hoặc reset trực tiếp DB nếu cần"],
    [6, "Books - List", "Kiểm tra API trả đúng phân trang", "GET {{baseUrl}}/api/books?page=1&limit=10", "SELECT COUNT(*) FROM Book WHERE isDeleted=0;\nKỳ vọng total khớp pagination.total", "Không cần rollback vì chỉ đọc"],
    [7, "Books - Search", "Kiểm tra search theo keyword", "GET {{baseUrl}}/api/books?search=clean", "SELECT * FROM Book b LEFT JOIN Author a ON b.authorId=a.id WHERE b.title LIKE '%clean%' OR b.isbn LIKE '%clean%' OR b.publisher LIKE '%clean%' OR b.description LIKE '%clean%' OR a.fullName LIKE '%clean%';", "Không cần rollback vì chỉ đọc"],
    [8, "Books - Filter EBOOK", "Kiểm tra chỉ trả sách có ebook", "GET {{baseUrl}}/api/books?availableAt=ebook", "SELECT DISTINCT bookId FROM BookEdition WHERE format='EBOOK' AND isDeleted=0;\nKỳ vọng book id trong response thuộc tập này", "Không cần rollback vì chỉ đọc"],
    [9, "Books - Filter Book Copy", "Kiểm tra chỉ trả sách có bản in sẵn", "GET {{baseUrl}}/api/books?availableAt=book-copy", "SELECT DISTINCT bookId FROM BookItem WHERE status='AVAILABLE' AND isDeleted=0;", "Không cần rollback vì chỉ đọc"],
    [10, "Books - Sort", "Kiểm tra thứ tự sắp xếp", "GET {{baseUrl}}/api/books?sortBy=publishYear&sortOrder=desc", "SELECT id,title,publishYear FROM Book WHERE isDeleted=0 ORDER BY publishYear DESC;\nKỳ vọng thứ tự response tương ứng", "Không cần rollback vì chỉ đọc"],
]

for r_idx, row in enumerate(rows, 3):
    for c_idx, val in enumerate(row, 1):
        cell = ws.cell(row=r_idx, column=c_idx, value=val)
        cell.alignment = center_align if c_idx == 1 else cell_align
        cell.border = thin_border
        if c_idx in (4, 5, 6):
            cell.fill = postman_fill

# Add variables section
start = len(rows) + 5
ws.merge_cells(start_row=start, start_column=1, end_row=start, end_column=6)
ws.cell(row=start, column=1, value="Biến môi trường Postman đề xuất").font = Font(bold=True)
ws.cell(row=start, column=1).fill = section_fill
for c in range(1, 7):
    ws.cell(row=start, column=c).border = thin_border
vars_rows = [
    ["baseUrl", base_url, "URL local server Next.js"],
    ["accessToken", "<token lấy từ API login>", "Dùng cho endpoint cần xác thực"],
    ["refreshToken", "<token lấy từ API login/cookie>", "Dùng cho refresh/logout"],
    ["testEmail", "nguyenvana_postman@example.com", "Email test nên dùng riêng để cleanup"],
]
for i, row in enumerate(vars_rows, start + 1):
    ws.cell(row=i, column=1, value=row[0]).border = thin_border
    ws.cell(row=i, column=2, value=row[1]).border = thin_border
    ws.cell(row=i, column=3, value=row[2]).border = thin_border
    for c in range(1, 4):
        ws.cell(row=i, column=c).alignment = cell_align

wb.save(path)
print(f"DONE: updated {path}")