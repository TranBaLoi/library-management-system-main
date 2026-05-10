# -*- coding: utf-8 -*-
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

path = r"C:\Users\Loi\Desktop\SQA\UnitTestReport_LibraryMS_Postman.xlsx"
wb = load_workbook(path)

thin_border = Border(left=Side(style="thin"), right=Side(style="thin"), top=Side(style="thin"), bottom=Side(style="thin"))
pass_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
postman_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
header_fill = PatternFill(start_color="D6E4F0", end_color="D6E4F0", fill_type="solid")
high_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
section_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
cell_align = Alignment(vertical="top", wrap_text=True)
center_align = Alignment(horizontal="center", vertical="top", wrap_text=True)

def style_cell(cell, center=False, fill=None, bold=False):
    cell.border = thin_border
    cell.alignment = center_align if center else cell_align
    if fill:
        cell.fill = fill
    if bold:
        cell.font = Font(bold=True)

# ===== Sheet 3: thêm 2 test case Auth =====
ws = wb["3.TC_XacThuc"]

# Xóa nếu đã tồn tại để tránh duplicate
ids_to_remove = {"TC_AUTH_REG_09", "TC_AUTH_CHANGEPWD_06"}
for row in range(ws.max_row, 2, -1):
    if ws.cell(row=row, column=1).value in ids_to_remove:
        ws.delete_rows(row, 1)

# Tìm vị trí chèn REG_09 sau REG_08
insert_reg_row = None
for row in range(3, ws.max_row + 1):
    if ws.cell(row=row, column=1).value == "TC_AUTH_REG_08":
        insert_reg_row = row + 1
        break

# Tìm vị trí chèn CHANGEPWD_06 sau CHANGEPWD_05 sau khi chèn reg
insert_changepwd_row = None

postman_col = None
for c in range(1, ws.max_column + 1):
    val = ws.cell(row=2, column=c).value
    if val and "Postman" in str(val):
        postman_col = c
        break
if not postman_col:
    postman_col = ws.max_column + 1
    ws.cell(row=2, column=postman_col, value="Lệnh Postman & CheckDB (Docker)")
    style_cell(ws.cell(row=2, column=postman_col), center=True, fill=header_fill, bold=True)
ws.column_dimensions[get_column_letter(postman_col)].width = 80

reg_row = [
    "TC_AUTH_REG_09",
    "Đăng ký vẫn thành công khi đồng bộ Gorse bị lỗi",
    "Dữ liệu đăng ký hợp lệ, nhưng GorseService.insertUser bị lỗi/reject",
    "User vẫn được tạo thành công, API không fail; chỉ ghi log lỗi Gorse",
    "PASS",
    "Bổ sung để cover nhánh catch ở auth.service.ts line 108; Gorse là external service nên lỗi không được làm hỏng đăng ký",
    "Unit test dùng mock GorseService.insertUser.mockRejectedValueOnce(new Error('Gorse down')).\n\nPostman khó giả lập Gorse down nếu không tắt service Gorse. Có thể kiểm thủ công bằng cách dừng container gorse rồi gọi POST /api/auth/register.\n\nCheckDB Docker:\ndocker exec -it library-mysql mysql -u root -pyour_mysql_root_password library_management -e \"SELECT id,email,role,status FROM User WHERE email='gorseerror@example.com';\"\nKỳ vọng: vẫn có user được tạo."
]

if insert_reg_row:
    ws.insert_rows(insert_reg_row, 1)
    for c, value in enumerate(reg_row, 1):
        cell = ws.cell(row=insert_reg_row, column=c, value=value)
        style_cell(cell, center=(c in (1, 5)), fill=(pass_fill if c == 5 else postman_fill if c == postman_col else None))

for row in range(3, ws.max_row + 1):
    if ws.cell(row=row, column=1).value == "TC_AUTH_CHANGEPWD_05":
        insert_changepwd_row = row + 1
        break

change_row = [
    "TC_AUTH_CHANGEPWD_06",
    "Ném lỗi khi mật khẩu mới quá yếu",
    "currentPassword='OldPass@123', newPassword='12345678', confirmNewPassword='12345678'",
    "Ném lỗi ValidationError chứa 'Password validation failed'; không gọi findUnique và không update DB",
    "PASS",
    "Bổ sung để cover nhánh validate mật khẩu mới yếu ở auth.service.ts line 272",
    "Postman:\nPOST http://localhost:3000/api/auth/change-password\nBody JSON:\n{\"currentPassword\":\"OldPass@123\",\"newPassword\":\"12345678\",\"confirmNewPassword\":\"12345678\"}\nKỳ vọng: HTTP 400.\n\nCheckDB Docker:\ndocker exec -it library-mysql mysql -u root -pyour_mysql_root_password library_management -e \"SELECT password FROM User WHERE email='nguyenvana_postman@example.com';\"\nKỳ vọng: password hash không đổi."
]

if insert_changepwd_row:
    ws.insert_rows(insert_changepwd_row, 1)
    for c, value in enumerate(change_row, 1):
        cell = ws.cell(row=insert_changepwd_row, column=c, value=value)
        style_cell(cell, center=(c in (1, 5)), fill=(pass_fill if c == 5 else postman_fill if c == postman_col else None))

# cập nhật title/note nếu cần
for r in range(3, ws.max_row + 1):
    if ws.cell(row=r, column=5).value == "PASS":
        ws.cell(row=r, column=5).fill = pass_fill
        ws.cell(row=r, column=5).alignment = center_align

# ===== Sheet 5: Execution Report =====
ws5 = wb["5.Ket qua thuc thi"]
# Cập nhật Auth Register từ 8 -> 9, ChangePwd từ 5 -> 6, tổng 54
for row in range(1, ws5.max_row + 1):
    suite = ws5.cell(row=row, column=1).value
    if suite and "TC_AUTH_REG" in str(suite):
        ws5.cell(row=row, column=2, value=9)
        ws5.cell(row=row, column=3, value=9)
        ws5.cell(row=row, column=4, value=0)
        ws5.cell(row=row, column=5, value="100%")
    if suite and "TC_AUTH_CHANGEPWD" in str(suite):
        ws5.cell(row=row, column=2, value=6)
        ws5.cell(row=row, column=3, value=6)
        ws5.cell(row=row, column=4, value=0)
        ws5.cell(row=row, column=5, value="100%")
    if suite and "TỔNG CỘNG" in str(suite):
        ws5.cell(row=row, column=2, value=54)
        ws5.cell(row=row, column=3, value=54)
        ws5.cell(row=row, column=4, value=0)
        ws5.cell(row=row, column=5, value="100%")

for row in range(1, ws5.max_row + 1):
    for col in range(1, min(ws5.max_column, 5) + 1):
        ws5.cell(row=row, column=col).border = thin_border
        ws5.cell(row=row, column=col).alignment = center_align if col > 1 else cell_align
        if col == 5 and ws5.cell(row=row, column=col).value == "100%":
            ws5.cell(row=row, column=col).fill = pass_fill

# Cập nhật dòng kết quả console nếu có
for row in range(1, ws5.max_row + 1):
    val = ws5.cell(row=row, column=1).value
    if val and "Kết quả" in str(val):
        ws5.cell(row=row, column=2, value="Test Suites: 2 passed, 2 total | Tests: 54 passed, 54 total | Time: 1.774s")
        ws5.cell(row=row, column=2).alignment = cell_align

# ===== Sheet 6: Code Coverage =====
ws6 = wb["6.Code Coverage"]
# Cập nhật các dòng coverage hiện có
coverage_updates = {
    "services/auth.service.ts": ("100%", "96.72%", "100%", "100%", "Đã cover thêm line 108 và 272; còn branch phụ do toán tử điều kiện optional"),
    "services/book.service.ts": ("93.84%", "90.47%", "100%", "93.65%", "Uncovered: line 113-114, 117-118"),
}
for row in range(1, ws6.max_row + 1):
    file_val = ws6.cell(row=row, column=1).value
    if file_val in coverage_updates:
        stmts, branches, funcs, lines, note = coverage_updates[file_val]
        for col, value in enumerate([stmts, branches, funcs, lines, note], 2):
            ws6.cell(row=row, column=col, value=value)
            ws6.cell(row=row, column=col).alignment = center_align if col <= 5 else cell_align
            ws6.cell(row=row, column=col).border = thin_border
            if col <= 5:
                ws6.cell(row=row, column=col).fill = high_fill

# Xóa hoặc cập nhật dòng tổng cũ nếu có
for row in range(1, ws6.max_row + 1):
    val = ws6.cell(row=row, column=1).value
    if val and "Tổng kết" in str(val):
        ws6.cell(row=row, column=2, value="97.45%")
        ws6.cell(row=row, column=3, value="93.10%")
        ws6.cell(row=row, column=4, value="100%")
        ws6.cell(row=row, column=5, value="97.41%")
        ws6.cell(row=row, column=6, value="Coverage sau khi giới hạn phạm vi vào 2 service chính và thêm 2 test case mới")
        for col in range(2, 6):
            ws6.cell(row=row, column=col).fill = high_fill

# Thêm ghi chú summary mới ở cuối nếu chưa có
summary_text = "Kết quả mới: All files = 97.45% statements, 93.10% branches, 100% functions, 97.41% lines; Tests = 54 passed / 54 total."
exists = False
for row in range(1, ws6.max_row + 1):
    if ws6.cell(row=row, column=1).value == "Ghi chú cập nhật 2026-05-09":
        exists = True
        ws6.cell(row=row, column=2, value=summary_text)
        break
if not exists:
    new_row = ws6.max_row + 2
    ws6.cell(row=new_row, column=1, value="Ghi chú cập nhật 2026-05-09")
    ws6.cell(row=new_row, column=2, value=summary_text)
    ws6.merge_cells(start_row=new_row, start_column=2, end_row=new_row, end_column=6)
    ws6.cell(row=new_row, column=1).font = Font(bold=True)
    ws6.cell(row=new_row, column=1).fill = section_fill
    ws6.cell(row=new_row, column=2).alignment = cell_align
    for col in range(1, 7):
        ws6.cell(row=new_row, column=col).border = thin_border

# ===== Sheet 7: Prompt / refs =====
ws7 = wb["7.Tai lieu Prompt"]
last = ws7.max_row + 2
ws7.cell(row=last, column=1, value="Cập nhật")
ws7.cell(row=last, column=2, value="Thêm 2 test case TC_AUTH_REG_09 và TC_AUTH_CHANGEPWD_06 để cover line 108 và 272 của auth.service.ts; cập nhật tổng test từ 52 lên 54 và coverage All files lên 97.45%.")
for col in range(1, 3):
    ws7.cell(row=last, column=col).border = thin_border
    ws7.cell(row=last, column=col).alignment = cell_align
ws7.cell(row=last, column=1).font = Font(bold=True)
ws7.cell(row=last, column=1).fill = section_fill

# Lưu bản mới và ghi đè file hiện tại
wb.save(path)
print(f"DONE: updated {path}")